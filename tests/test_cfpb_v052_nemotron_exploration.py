from __future__ import annotations

import ast
import base64
import gzip
import hashlib
import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from scripts.build_cfpb_v052_nemotron_exploration_notebook import ROOT, build_notebook, inputs, digest
from scripts.generation.nemo_data_designer import cfpb_v052_nemotron_exploration_v01 as w


@pytest.fixture
def data(tmp_path):
    seed = {"seed_id": "seed_1", "product": "Card", "sub_product": "Credit card",
            "issue": "Billing", "sub_issue": "Wrong amount",
            "generation_grounding_excerpt": "The customer noticed an unexpected amount on the card statement and asked the institution to explain the transaction."}
    row = {"seed_id": "seed_1", "conversation_length": 4, "dialogue": {
        "conversation": [{"role": "user", "content": "I have a question about a charge."},
                         {"role": "assistant", "content": "What do you remember about it?"},
                         {"role": "user", "content": "The amount looks different."},
                         {"role": "assistant", "content": "You can ask through the institution's official channel."}],
        "synthetic_case_summary": "A customer asks about a charge.", "privacy_notes": ["Generic details."]}}
    prepared = tmp_path / "prepared.jsonl"
    raw = tmp_path / "raw.jsonl"
    prepared.write_text(json.dumps(seed) + "\n", encoding="utf-8")
    raw.write_text(json.dumps(row) + "\n", encoding="utf-8")
    return prepared, raw, seed, row


def test_write_once_preserves_existing(tmp_path):
    p = tmp_path / "record.json"
    w.save_json(p, {"a": 1})
    w.save_json(p, {"a": 1})
    with pytest.raises(ValueError, match="Existing artifact differs"):
        w.save_json(p, {"a": 2})


@pytest.mark.parametrize("mode", ["missing", "duplicate", "unknown"])
def test_raw_coverage_fails_closed(data, mode):
    prepared, raw, _, row = data
    rows = [] if mode == "missing" else [row, row] if mode == "duplicate" else [{**row, "seed_id": "other"}]
    raw.write_text("\n".join(json.dumps(r) for r in rows), encoding="utf-8")
    with pytest.raises(ValueError):
        w.check_coverage(raw, prepared)


def test_review_pack_pending_formula_safe_html_escaped(data, tmp_path):
    prepared, raw, seed, _ = data
    seed["product"] = "=HYPERLINK(\"https://invalid.test\")"
    seed["generation_grounding_excerpt"] += " <script>alert('x')</script>"
    prepared.write_text(json.dumps(seed), encoding="utf-8")
    before = {p: w.sha256_file(p) for p in [raw, prepared]}
    book = w.build_review_pack(raw, prepared, tmp_path / "review")
    wb = load_workbook(book)
    sheet = wb["Review"]
    assert sheet.freeze_panes == "F2"
    assert sheet.auto_filter.ref
    assert sheet["B2"].data_type == "s"
    assert sheet["F2"].alignment.wrap_text
    assert len(sheet.data_validations.dataValidation) == 7
    output = (book.parent / "review_readable.html").read_text(encoding="utf-8")
    assert "<script>" not in output and "&lt;script&gt;" in output
    report = w.summarize_review(book)
    assert report["pending_rows"] == 1 and report["reviewed_rows"] == 0
    assert report["decisions_by_reviewer_kind"] == {"human": {}, "llm": {}}
    assert report["policy"]["benchmark_eligible"] is False
    technical = json.loads((book.parent / "technical_check_report.json").read_text())
    assert technical["auto_accepted_rows"] == 0
    assert before == {p: w.sha256_file(p) for p in [raw, prepared]}


def fill_review(book, kind="human", decision="usable"):
    wb = load_workbook(book)
    sheet = wb["Review"]
    values = {**{key: "pass" for key in w.DIMENSIONS}, "review_decision": decision,
              "reviewer_kind": kind, "reviewer_id": "test-reviewer", "reviewed_utc": "2026-09-19T12:00:00Z",
              "review_status": "reviewed"}
    for key, value in values.items():
        sheet.cell(2, w.COLUMNS.index(key) + 1, value)
    wb.save(book)


def test_review_edits_preserved_and_llm_not_human(data, tmp_path):
    prepared, raw, _, _ = data
    book = w.build_review_pack(raw, prepared, tmp_path / "review")
    fill_review(book, kind="llm")
    reviewed_hash = w.sha256_file(book)
    w.build_review_pack(raw, prepared, book.parent)
    assert w.sha256_file(book) == reviewed_hash
    result = w.summarize_review(book)
    assert result["decisions_by_reviewer_kind"] == {"human": {}, "llm": {"usable": 1}}
    assert result["pending_rows"] == 0


def test_summary_refuses_changed_evidence(data, tmp_path):
    prepared, raw, _, _ = data
    book = w.build_review_pack(raw, prepared, tmp_path / "review")
    wb = load_workbook(book)
    wb["Review"]["F2"] = "changed"
    wb.save(book)
    with pytest.raises(ValueError, match="Reference evidence changed"):
        w.summarize_review(book)


@pytest.mark.parametrize("key,value", [("reviewer_id", ""), ("reviewed_utc", "2026-09-19T12:00:00"),
                                       ("factual_fidelity", "fail"), ("reviewer_kind", "robot")])
def test_summary_rejects_incomplete_or_inconsistent_review(data, tmp_path, key, value):
    prepared, raw, _, _ = data
    book = w.build_review_pack(raw, prepared, tmp_path / "review")
    fill_review(book)
    wb = load_workbook(book)
    wb["Review"].cell(2, w.COLUMNS.index(key) + 1, value)
    wb.save(book)
    with pytest.raises(ValueError):
        w.summarize_review(book)


def test_no_automatic_retry_after_failed_generation(data, tmp_path):
    prepared, _, _, _ = data
    assert w.resolve_cached_raw(tmp_path, prepared) is None
    w.save_json(tmp_path / "generation_started.json", {"started": True})
    with pytest.raises(RuntimeError, match="Previous API attempt"):
        w.resolve_cached_raw(tmp_path, prepared)


def test_recover_completed_batch_and_detect_tampering(data, tmp_path):
    prepared, _, _, row = data
    raw = tmp_path / "raw/data_designer/dataset/parquet-files/batch_00000.parquet"
    raw.parent.mkdir(parents=True)
    pd.DataFrame([row]).to_parquet(raw)
    w.save_json(tmp_path / "generation_started.json", {"started": True})
    assert w.resolve_cached_raw(tmp_path, prepared).resolve() == raw.resolve()
    assert w.resolve_cached_raw(tmp_path, prepared).resolve() == raw.resolve()
    prepared.write_text(prepared.read_text() + "\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Prepared input differs"):
        w.resolve_cached_raw(tmp_path, prepared)


def test_malformed_dialogue_still_in_reading_pack(data, tmp_path):
    prepared, raw, _, row = data
    row["dialogue"] = "not valid json"
    raw.write_text(json.dumps(row), encoding="utf-8")
    records = w.review_records(raw, prepared)
    assert len(records) == 1 and "not valid json" in records[0]["generated_dialogue"]
    assert "conversation_schema_invalid" in records[0]["deterministic_flags_json"]


def test_workbook_never_silently_truncates(tmp_path):
    record = {key: "" for key in w.COLUMNS}
    record["generated_dialogue"] = "x" * 32768
    with pytest.raises(ValueError, match="cell limit"):
        w.make_workbook(tmp_path / "review.xlsx", [record], {})


def test_notebook_cells_compile_and_payload_hashes_match():
    notebook, inventory = build_notebook()
    codes = [c for c in notebook["cells"] if c["cell_type"] == "code"]
    assert len(codes) == 8 and len(inventory) == 11
    for c in codes:
        ast.parse(c["source"])
        assert not c["outputs"] and c["execution_count"] is None
    payload = next(n.value for n in ast.parse(codes[2]["source"]).body
                   if isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "EMBEDDED" for t in n.targets))
    for item in json.loads(ast.literal_eval(payload.args[0])).values():
        raw = gzip.decompress(base64.b64decode(item["data"]))
        assert hashlib.sha256(raw).hexdigest() == item["sha256"]
        ast.parse(raw)
    assert "RUN_GENERATION = False" in codes[5]["source"]
    assert "preview(" not in codes[5]["source"]


def test_real_inputs_prepare_twenty_disjoint_rows_offline(tmp_path):
    inventory = {p.as_posix(): digest(ROOT / p) for p in inputs()}
    sample, manifest, info = w.prepare_run(ROOT, tmp_path, inventory)
    rows = w.load_records(sample)
    assert len(rows) == 20
    assert info["prior_sample_exclusion"]["selected_overlap"] == []
    assert len(info["prior_sample_exclusion"]["seed_ids"]) == 20
    assert info["offset_invariant"]["full_text_prefix_rows_verified"] == 3004
    assert info["offset_invariant"]["finding_offsets_verified"] == 780
    assert all(r["offset_invariant_verified"] and not r["benchmark_eligible"] for r in rows)
    assert all("seed_text" not in r and "seed_narrative_excerpt" not in r for r in rows)
    before = w.sha256_file(manifest)
    sample2, _, _ = w.prepare_run(ROOT, tmp_path, inventory)
    assert sample2 == sample and w.sha256_file(manifest) == before
    other, _, _ = w.prepare_run(ROOT, tmp_path / "repeat", inventory)
    assert w.sha256_file(other) == w.sha256_file(sample)
