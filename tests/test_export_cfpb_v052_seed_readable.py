import importlib.util
import json
import re
import shutil
import subprocess
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("seed_readable", ROOT / "scripts/export_cfpb_v052_seed_readable.py")
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)


def test_browser_text_is_safe_and_preserved():
    text = '</script><script>alert("bad")</script> & <img src=x onerror=alert(1)>'
    records = [{"seed_id": "a", "seed_text": text, "missing": None, "unicode": "\u2028"}]
    rendered = m.render_browser("<bad title>", records)
    data = re.search(r'<script id="data" type="application/json">(.*?)</script>', rendered, re.S).group(1)
    assert json.loads(data) == records
    assert '<script>alert' not in rendered
    assert "<bad title>" not in rendered
    assert 'textContent' in rendered and 'fetch(' not in rendered and 'innerHTML' not in rendered


def test_long_text_and_formula_like_text_are_not_lost(tmp_path):
    long = "abcdefghij" * 5000
    records = [{"seed_id": "a", "seed_text": long, "product": "=1+1", "control": "a\x01b"}]
    path = tmp_path / "test.xlsx"
    report = m.export_excel(path, records, list(records[0]), "test", ["source.parquet"])
    wb = load_workbook(path)
    assert report["long_text_parts"] == 2
    assert report["escaped_control_characters"] == 1
    rows = list(wb["Long_text"].values)
    assert "".join(r[-1] for r in rows[1:]) == long
    all_fields = wb["All_fields"]
    assert all_fields["D2"].value == "=1+1" and all_fields["D2"].data_type == "s"
    assert all_fields["E2"].value == "a\\u0001b"
    assert all(s.freeze_panes == "B2" and s.auto_filter.ref for s in wb)


def test_pairs_are_compared_and_differences_fail_closed(tmp_path):
    path = tmp_path / "input.jsonl"
    frame = pd.DataFrame({"seed_id": ["a"], "time": [pd.Timestamp("2025-01-01", tz="UTC")]})
    path.write_text(json.dumps(m.source_records(frame, iso=False)[0]) + "\n", encoding="utf-8")
    m.verify_jsonl_pair(frame, path)
    path.write_text('{"seed_id":"different"}\n', encoding="utf-8")
    with pytest.raises(ValueError, match="differ"):
        m.verify_jsonl_pair(frame, path)


def test_existing_output_and_outside_source_refused(tmp_path):
    out = tmp_path / "existing"
    out.mkdir()
    with pytest.raises(FileExistsError):
        m.export_all(tmp_path, out)
    with pytest.raises(ValueError):
        m.export_all(tmp_path, tmp_path.parent / "outside")


def test_javascript_syntax():
    node = shutil.which("node")
    if not node:
        pytest.skip("Node not available")
    script = re.search(r"<script>\n(.*?)</script>", m.BROWSER, re.S).group(1)
    result = subprocess.run([node, "--check"], input=script, encoding="utf-8", capture_output=True)
    assert result.returncode == 0, result.stderr


def test_completed_exports_cover_all_sources_and_html_rows():
    out = m.SOURCE / "readable_v01"
    if not (out / "export_manifest.json").exists():
        pytest.skip("Full export has not finished")
    manifest = json.loads((out / "export_manifest.json").read_text(encoding="utf-8"))
    assert manifest["source_files_unchanged"] is True
    assert len(manifest["datasets"]) == 7
    covered = {f for d in manifest["datasets"] for f in d["sources"]}
    assert covered == {p.name for p in m.SOURCE.iterdir() if p.suffix in {".parquet", ".jsonl"}}
    for name, expected in manifest["original_files_sha256"].items():
        assert m.sha(m.SOURCE / name) == expected
    for name, expected in manifest["outputs"].items():
        assert m.sha(out / name) == expected
    for d in manifest["datasets"]:
        page = (out / (d["name"] + "_browser.html")).read_text(encoding="utf-8")
        data = re.search(r'<script id="data" type="application/json">(.*?)</script>', page, re.S).group(1)
        records = json.loads(data)
        assert records == m.source_records(pd.read_parquet(m.SOURCE / (d["name"] + ".parquet")))
        assert len(records) == d["rows"]
        wb = load_workbook(out / (d["name"] + "_readable.xlsx"), read_only=True)
        assert sum(1 for _ in wb["All_fields"].values) - 1 == d["rows"]
        wb.close()
