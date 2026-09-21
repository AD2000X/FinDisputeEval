import ast
import base64
import gzip
import hashlib
import importlib.util
import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def module(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


m = module("language_comparison", "scripts/evaluation/cfpb_v052_linguistic_comparison_v01.py")
b = module("language_builder", "scripts/build_cfpb_v052_linguistic_comparison_notebook.py")


def test_normalization_and_short_mattr():
    assert m.normalize("Ｉ can’t") == "I can't"
    assert np.isnan(m.mattr(["a"] * 24))
    assert m.mattr(["a"] * 25) == pytest.approx(1 / 25)
    assert m.mattr([str(i) for i in range(26)]) == 1


def test_no_cross_sentence_or_turn_ngrams():
    rows = [{"seed_id": key, "stage": "generated_assistant", "tokens": words}
            for key in ["a", "b"] for words in [["please", "call"], ["your", "bank"]]]
    assert m.repeated_phrases(rows).empty
    rows += [{"seed_id": key, "stage": "generated_assistant", "tokens": ["please", "call", "the", "bank"]}
             for key in ["a", "b"]]
    frame = m.repeated_phrases(rows)
    assert set(frame.distinct_cases) == {2}
    assert "please call the bank" in set(frame.phrase)


def test_repetition_requires_distinct_cases():
    row = {"seed_id": "a", "stage": "generated_user", "tokens": ["i", "called", "bank"]}
    assert m.repeated_phrases([row] * 10).empty


def test_question_pairs_do_not_assign_quality():
    cases = [{"seed_id": "a", "messages": [{"role": "assistant", "content": "When?"},
              {"role": "user", "content": "Yesterday."}, {"role": "assistant", "content": "Anything else?"}]}]
    pairs = m.interaction_pairs(cases)
    assert list(pairs.following_user_turn_present) == [True, False]
    assert set(pairs.adequacy) == {"NOT_ASSESSED"}


def test_duplicate_keys_fail_closed():
    with pytest.raises(ValueError, match="duplicate"):
        m.unique_index([{"seed_id": "a"}, {"seed_id": "a"}], "test")
    with pytest.raises(ValueError):
        m.unique_index([{"seed_id": None}], "test")


def test_path_escape_rejected(tmp_path):
    with pytest.raises(ValueError):
        m.resolve_inside(tmp_path, "../outside")
    with pytest.raises(ValueError):
        m.run_analysis(ROOT, ROOT / m.RUN_REL, "unused")


def test_workbook_safe_text_and_readability(tmp_path):
    path = tmp_path / "report.xlsx"
    m.export_workbook(path, {"Evidence": pd.DataFrame({"sentence": ["=1+1", "<script>text</script>", None]})})
    ws = load_workbook(path)["Evidence"]
    assert ws["A2"].data_type == "s"
    assert ws["A2"].value == "=1+1"
    assert ws["A2"].alignment.wrap_text
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref == "A1:A4"


def test_workbook_refuses_silent_truncation(tmp_path):
    with pytest.raises(ValueError, match="truncated"):
        m.export_workbook(tmp_path / "bad.xlsx", {"Evidence": pd.DataFrame({"text": ["a" * 32768]})})


def test_legacy_contract_matches_source():
    contract = b.legacy_contract()
    assert len(contract["patterns"]) == 7
    assert contract == json.loads(b.CONTRACT.read_text(encoding="utf-8"))
    import re
    pattern = re.compile(**contract["patterns"]["lx_negation_candidate"])
    assert pattern.search("can't")
    assert not pattern.search("can’t")  # Intentional unchanged legacy behavior.
    assert pattern.search(m.normalize("can’t"))


def test_notebook_valid_embedded_source_and_no_execution():
    notebook = json.loads(b.OUTPUT.read_text(encoding="utf-8"))
    built, _ = b.build_notebook()
    # zlib can emit different valid gzip streams across Python builds; compare
    # decoded bytes below, and compare all non-payload cells directly.
    assert len(notebook["cells"]) == len(built["cells"])
    for saved, generated in zip(notebook["cells"], built["cells"]):
        if "EMBEDDED = " not in saved["source"]:
            assert saved == generated
    found = False
    for cell in notebook["cells"]:
        if cell["cell_type"] != "code":
            continue
        assert cell["execution_count"] is None and cell["outputs"] == []
        tree = ast.parse(cell["source"])
        for node in tree.body:
            if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "EMBEDDED" for t in node.targets):
                embedded = ast.literal_eval(node.value)
                for name, entry in embedded.items():
                    payload = gzip.decompress(base64.b64decode(entry["gzip_base64"]))
                    assert hashlib.sha256(payload).hexdigest() == entry["sha256"]
                    assert payload == (b.SCRIPT if name.endswith(".py") else b.CONTRACT).read_bytes()
                found = True
    assert found


def test_real_input_lineage_and_all_20_cases_retained():
    inventory, raw, prepared = m.verify_inputs(ROOT)
    cases = m.assemble_cases(ROOT, raw, prepared)
    assert len(inventory) == 33
    assert len(cases) == 20
    assert sum(c["format_valid"] for c in cases) == 19
    assert any(c["actual_messages"] == 9 for c in cases)
    for c in cases:
        assert sum(len(c["texts"][s]) for s in ["generated_user", "generated_assistant"]) == c["actual_messages"]


def test_real_parser_separates_roles_and_excludes_placeholders():
    pytest.importorskip("spacy")
    nlp = m.load_parser()
    texts = {s: ["I did not authorize this. [REDACTED]"] for s in m.STAGES}
    texts["generated_user"] = ["Yes.", "No."]
    texts["generated_assistant"] = ["Please call your bank."]
    cases = [{"seed_id": "test", "release_split": "test", "product": "card", "issue": "test", "mood": "test", "texts": texts}]
    features, _, _, lemmas, _ = m.extract_features(cases, nlp, b.legacy_contract())
    source = features[features.stage == "grounding"].iloc[0]
    user = features[features.stage == "generated_user"].iloc[0]
    assert source.placeholder_count == 1
    assert source.word_count == 5
    assert source.negation_markers == 1
    assert user.units == 2 and user.word_count == 2
    assert np.isnan(user.mattr25_within_unit)
    assert "bank" not in lemmas[("test", "generated_user")]
    assert "bank" in lemmas[("test", "generated_assistant")]
