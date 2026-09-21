"""Read-only source/run analysis; exploratory linguistics, not an acceptance gate.

spaCy parses each turn separately. pandas/PyArrow, sklearn and openpyxl handle
tables, lexical statistics and exports. No generation or judge SDK is imported.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import importlib.metadata
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.feature_extraction.text import CountVectorizer, TfidfVectorizer

RUN_REL = "outputs/generation/smoke_only/cfpb_seed_v052_pipeline_override/exploratory_nemotron_super_v02/run_20260919T220844338002Z"
EDA_REL = "outputs/data_pipeline/cfpb_seed_source_eda/eda_v051/run_20260713T145423Z"
SEED_REL = "dataset/curated/seed_pools/cfpb_dispute/seed_v052"
EDA_TABLES = ["linguistic_summary.csv", "prevalence_sensitivity.csv", "tfidf_top_terms.csv",
              "frequent_ngrams.csv", "feature_registry.json"]
STAGES = ["seed_full", "source_excerpt", "grounding", "generated_user", "generated_assistant"]
PLACEHOLDERS = re.compile(r"\[(?:REDACTED|PII_CANDIDATE|DATE|AMOUNT)\]")
POLICY = {"pipeline_smoke_only": True, "privacy_verified": False, "benchmark_eligible": False,
          "distribution_calibration_eligible": False, "formal_pilot_allowed": False,
          "automatic_quality_decisions": False, "judge_used": False}
CONFIG = {"analysis_version": "v01", "spacy_model": "en_core_web_sm", "model_version": "3.8.0",
          "spacy_version": "3.8.7", "mattr_window": 25, "phrase_ngram_range": [3, 5],
          "phrase_min_distinct_cases": 2, "statistical_tests": False,
          "unit": "seed-stage; turns parsed separately", "normalization": "NFKC + straight apostrophes; placeholders excluded"}
CAUTIONS = [
    "探索性 smoke，非正式分布校準、品質通過判定或可發布 benchmark。",
    "20 筆為 coverage-oriented 抽樣；不代表母體。每案例權重相同；另按 release_split 分組。",
    "full seed→excerpt→grounding 分開；user 與 assistant 分開；不跨 turn 解析或建構片語。",
    "legacy 背景使用未改動的 v05.1 偵測器；normalized 指標使用統一新前處理，兩者不可混稱同一數值。",
    "EDA 的 raw、exact-deduplicated、family-weighted 分布分開保存；enrichment 不混入 population。",
    "來源 TF-IDF 與本批 TF-IDF 的語料及 IDF 不同，不能直接比較分數大小。",
    "spaCy 與 regex 都是候選測量工具；否定詞頻、詞彙相似度不證明授權狀態或事實一致。",
    "repair/emotion/authorization 等欄位是候選詞彙命中，不是經驗證的語用或心理狀態標籤。",
    "問句與下一輪回應只提供並排例句；不自動判斷澄清是否有效、回答是否充分。",
    "mood、訊息長度與匿名化由 prompt 控制；相關差異不是自然母體差異。",
    "保留全部 20 筆，包括格式異常的 9 則對話，不只分析通過規則的資料。",
    "不分析語音/韻律；不自動評定法律、隱私、語意正確性；不計算 p 值或通過門檻。",
]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def read_jsonl(path):
    return [json.loads(line) for line in Path(path).read_text(encoding="utf-8").splitlines() if line.strip()]


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def resolve_inside(root, relative):
    path = (Path(root) / relative).resolve()
    if Path(root).resolve() not in path.parents:
        raise ValueError(f"Path escapes manifest root: {relative}")
    return path


def verify_inputs(root, run_rel=RUN_REL):
    root = Path(root).resolve()
    run, eda, seed = root / run_rel, root / EDA_REL, root / SEED_REL
    paths = [run / "evidence_manifest.json", eda / "manifest.json", seed / "seed_v052_manifest.json"]
    evidence = read_json(paths[0])
    for relative, expected in evidence["files"].items():
        path = resolve_inside(run, relative)
        if sha(path) != expected:
            raise ValueError(f"Run evidence hash mismatch: {relative}")
        paths.append(path)
    eda_manifest = read_json(eda / "manifest.json")
    for name in EDA_TABLES:
        path = eda / name
        if sha(path) != eda_manifest["outputs"][name]["sha256"]:
            raise ValueError(f"EDA hash mismatch: {name}")
        paths.append(path)
    seed_manifest = read_json(seed / "seed_v052_manifest.json")
    for key in ["seed_parquet", "generation_jsonl"]:
        item = seed_manifest["outputs"][key]
        path = resolve_inside(seed, item["path"])
        if sha(path) != item["sha256"]:
            raise ValueError(f"Seed hash mismatch: {key}")
        paths.append(path)
    prepared_manifest = read_json(run / "prepared_inputs/pipeline_smoke_input_manifest.json")
    if prepared_manifest["source_seed_manifest_sha256"] != sha(seed / "seed_v052_manifest.json"):
        raise ValueError("Prepared input refers to a different seed release")
    pointer = read_json(run / "raw_output_pointer.json")
    raw = resolve_inside(run, pointer["path"])
    prepared = run / "prepared_inputs/nemo_seed_v052_pipeline_smoke_20.jsonl"
    if sha(raw) != pointer["sha256"] or sha(prepared) != pointer["prepared_sha256"]:
        raise ValueError("Raw/prepared pointer binding mismatch")
    return {p.relative_to(root).as_posix(): sha(p) for p in sorted(set(paths))}, raw, prepared


def unique_index(rows, label):
    keys = [row.get("seed_id") for row in rows]
    if any(not isinstance(k, str) or not k for k in keys) or len(set(keys)) != len(keys):
        raise ValueError(f"Missing/duplicate seed IDs: {label}")
    return dict(zip(keys, rows))


def assemble_cases(root, raw, prepared):
    root = Path(root)
    seeds = unique_index(read_jsonl(prepared), "prepared")
    generated = unique_index(pd.read_parquet(raw).to_dict("records"), "raw")
    if len(seeds) != 20 or set(seeds) != set(generated):
        raise ValueError("Expected exactly 20 matching raw/prepared seeds")
    full = unique_index(pd.read_parquet(root / SEED_REL / "cfpb_seed_v052.parquet",
                                       columns=["seed_id", "seed_text"]).to_dict("records"), "seed")
    excerpts = unique_index(read_jsonl(root / SEED_REL / "cfpb_seed_v052_generation_input.jsonl"), "excerpt")
    result = []
    for key, seed in seeds.items():
        text = full[key]["seed_text"]
        excerpt = excerpts[key]["seed_narrative_excerpt"]
        if not text.startswith(excerpt):
            raise ValueError("Full seed/excerpt prefix invariant violated")
        if hashlib.sha256(excerpt.encode()).hexdigest() != seed["source_excerpt_sha256"]:
            raise ValueError("Excerpt SHA differs from prepared source")
        dialogue = generated[key]["dialogue"]
        if isinstance(dialogue, str):
            dialogue = json.loads(dialogue)
        messages = list(dialogue.get("conversation", []))
        if not messages or any(not isinstance(m, dict) or m.get("role") not in {"user", "assistant"}
                               or not isinstance(m.get("content"), str) for m in messages):
            raise ValueError(f"Unparseable messages: {key}; cannot silently omit a case")
        requested = int(generated[key]["conversation_length"])
        valid = (len(messages) == requested and requested in {4, 6, 8}
                 and all(m["role"] == ("user" if i % 2 == 0 else "assistant") for i, m in enumerate(messages)))
        result.append({"seed_id": key, "release_split": seed["release_split"], "product": seed["product"],
                       "issue": seed["issue"], "mood": generated[key].get("customer_mood", ""),
                       "expected_messages": requested, "actual_messages": len(messages), "format_valid": valid,
                       "messages": messages, "texts": {
                           "seed_full": [text], "source_excerpt": [excerpt],
                           "grounding": [seed["generation_grounding_excerpt"]],
                           "generated_user": [m["content"] for m in messages if m["role"] == "user"],
                           "generated_assistant": [m["content"] for m in messages if m["role"] == "assistant"]}})
    return result


def normalize(text):
    return unicodedata.normalize("NFKC", text).translate(str.maketrans({"’": "'", "‘": "'"}))


def load_parser():
    import spacy
    from spacy.symbols import ORTH
    if spacy.__version__ != CONFIG["spacy_version"]:
        raise RuntimeError("Install pinned spaCy " + CONFIG["spacy_version"] + "; no silent fallback")
    nlp = spacy.load(CONFIG["spacy_model"], disable=["ner"])
    if nlp.meta["version"] != CONFIG["model_version"] or not {"parser", "tagger", "lemmatizer"}.issubset(nlp.pipe_names):
        raise RuntimeError("Required parser/tagger/lemmatizer/model version missing")
    for placeholder in ["[REDACTED]", "[PII_CANDIDATE]", "[DATE]", "[AMOUNT]"]:
        nlp.tokenizer.add_special_case(placeholder, [{ORTH: placeholder}])
    return nlp


def mattr(tokens, window=25):
    if len(tokens) < window:
        return np.nan
    return float(np.mean([len(set(tokens[i:i + window])) / window for i in range(len(tokens) - window + 1)]))


def lexical_token(token):
    return not (token.is_space or token.is_punct or PLACEHOLDERS.fullmatch(token.text)) and any(c.isalnum() for c in token.text)


def safe_rate(count, denominator, scale=100):
    return count / denominator * scale if denominator else np.nan


def extract_features(cases, nlp, contract):
    patterns = {name: re.compile(item["pattern"], item["flags"]) for name, item in contract["patterns"].items()}
    units = [(case["seed_id"], stage, i, text) for case in cases for stage in STAGES
             for i, text in enumerate(case["texts"][stage])]
    docs = list(nlp.pipe((normalize(u[3]) for u in units), batch_size=64))
    by_stage = defaultdict(list)
    by_doc = {}
    for unit, doc in zip(units, docs):
        by_stage[unit[:2]].append((unit[3], doc))
        by_doc[unit[:3]] = doc
    records, sentence_rows, cue_rows = [], [], []
    dictionaries = {}
    for case in cases:
        for stage in STAGES:
            items = by_stage[(case["seed_id"], stage)]
            tokens = [t for _, doc in items for t in doc if lexical_token(t)]
            lemmas = [t.lemma_.lower() for t in tokens if t.pos_ in {"NOUN", "PROPN", "VERB", "ADJ", "ADV"} and not t.is_stop]
            dictionaries[(case["seed_id"], stage)] = lemmas
            word_count = len(tokens)
            sentences = [s for _, doc in items for s in doc.sents if any(lexical_token(t) for t in s)]
            question_count = sum("?" in s.text for s in sentences)
            passive = sum(t.dep_ in {"nsubjpass", "auxpass"} or "Pass" in t.morph.get("Voice") for t in tokens)
            negations = sum(t.dep_ == "neg" for t in tokens)
            record = {"seed_id": case["seed_id"], "stage": stage, "release_split": case["release_split"],
                      "product": case["product"], "issue": case["issue"], "mood": case["mood"],
                      "units": len(items), "word_count": word_count, "sentence_count": len(sentences),
                      "words_per_sentence": safe_rate(word_count, len(sentences), 1),
                      "words_per_unit": safe_rate(word_count, len(items), 1),
                      "placeholder_count": sum(len(PLACEHOLDERS.findall(doc.text)) for _, doc in items),
                      "negation_markers": negations, "negation_per100": safe_rate(negations, word_count),
                      "passive_markers_per100": safe_rate(passive, word_count),
                      "agent_markers_per100": safe_rate(sum(t.dep_ == "agent" for t in tokens), word_count),
                      "pronouns_per100": safe_rate(sum(t.pos_ == "PRON" for t in tokens), word_count),
                      "past_verbs_per100": safe_rate(sum(t.pos_ in {"VERB", "AUX"} and "Past" in t.morph.get("Tense") for t in tokens), word_count),
                      "modal_tokens_per100": safe_rate(sum(t.tag_ == "MD" for t in tokens), word_count),
                      "clause_markers_per_sentence": safe_rate(sum(t.dep_ in {"advcl", "ccomp", "xcomp", "acl", "relcl"} for t in tokens), len(sentences), 1),
                      "content_pos_percent": safe_rate(sum(t.pos_ in {"NOUN", "PROPN", "VERB", "ADJ", "ADV"} for t in tokens), word_count),
                      "question_sentences": question_count, "question_sentence_percent": safe_rate(question_count, len(sentences)),
                      "mattr25_within_unit": np.nan,
                      "uppercase_letter_percent": safe_rate(sum(c.isupper() for _, d in items for c in PLACEHOLDERS.sub("", d.text) if c.isalpha()),
                                                              sum(c.isalpha() for _, d in items for c in PLACEHOLDERS.sub("", d.text)))}
            # Do not concatenate turns into fake continuous syntax or MATTR windows.
            diversity = [mattr([t.text.lower() for t in doc if lexical_token(t)]) for _, doc in items]
            defined = [v for v in diversity if not np.isnan(v)]
            record["mattr25_within_unit"] = float(np.mean(defined)) if defined else np.nan
            record["mattr_eligible_units"] = len(defined)
            for feature, pattern in patterns.items():
                record["legacy_" + feature] = int(any(pattern.search(original) for original, _ in items))
                count = sum(len(list(pattern.finditer(doc.text))) for _, doc in items)
                record["normalized_" + feature + "_per100"] = safe_rate(count, word_count)
            records.append(record)
            if stage in {"grounding", "generated_user", "generated_assistant"}:
                for i, (_, doc) in enumerate(items):
                    for s in doc.sents:
                        words = [t.text.lower() for t in s if lexical_token(t)]
                        sentence_rows.append({"seed_id": case["seed_id"], "stage": stage,
                                              "unit_index": i, "tokens": words})
                        cues = [name for name, p in patterns.items() if p.search(s.text)]
                        if cues or any(t.dep_ == "neg" for t in s):
                            cue_rows.append({"seed_id": case["seed_id"], "stage": stage, "unit_index": i,
                                             "sentence": s.text, "lexical_candidate_codes": ", ".join(cues),
                                             "negation_heads": ", ".join(f"{t.text} -> {t.head.text}" for t in s if t.dep_ == "neg"),
                                             "interpretation": "candidate evidence only; scope/meaning not adjudicated"})
    return pd.DataFrame(records), sentence_rows, pd.DataFrame(cue_rows), dictionaries, by_doc


def overview(features):
    numeric = [c for c in features.select_dtypes(include="number") if not c.startswith("legacy_")]
    output = []
    for split in ["ALL_EXPLORATORY"] + sorted(features.release_split.unique()):
        selected = features if split == "ALL_EXPLORATORY" else features[features.release_split == split]
        for stage, rows in selected.groupby("stage", sort=False):
            for feature in numeric:
                values = rows[feature].dropna()
                output.append({"release_split": split, "stage": stage, "feature": feature,
                               "cases": len(rows), "defined_cases": len(values),
                               "mean": values.mean(), "median": values.median(),
                               "p25": values.quantile(.25), "p75": values.quantile(.75)})
    return pd.DataFrame(output)


def paired(features, dictionaries):
    chosen = ["word_count", "words_per_sentence", "negation_per100", "modal_tokens_per100",
              "pronouns_per100", "past_verbs_per100", "mattr25_within_unit", "content_pos_percent"]
    rows = []
    for key, case in features.groupby("seed_id", sort=False):
        idx = case.set_index("stage")
        for a, b in [("seed_full", "source_excerpt"), ("source_excerpt", "grounding"),
                     ("grounding", "generated_user")]:
            for feature in chosen:
                rows.append({"seed_id": key, "release_split": idx.iloc[0]["release_split"],
                             "from_stage": a, "to_stage": b, "feature": feature,
                             "before": idx.loc[a, feature], "after": idx.loc[b, feature],
                             "delta": idx.loc[b, feature] - idx.loc[a, feature]})
    frame = pd.DataFrame(rows)
    ids = list(features.seed_id.drop_duplicates())
    texts = [" ".join(dictionaries[(key, stage)]) for stage in ["grounding", "generated_user"] for key in ids]
    v = TfidfVectorizer(token_pattern=r"(?u)\b\w[\w'-]*\b", lowercase=False)
    try:
        matrix = v.fit_transform(texts)
        cosines = np.asarray(matrix[:len(ids)].multiply(matrix[len(ids):]).sum(axis=1)).ravel()
    except ValueError as exc:
        if "empty vocabulary" not in str(exc):
            raise
        matrix, cosines = None, [np.nan] * len(ids)
    lexical = []
    for key, cosine in zip(ids, cosines):
        source, user = set(dictionaries[(key, "grounding")]), set(dictionaries[(key, "generated_user")])
        lexical.append({"seed_id": key, "grounding_content_lemma_types": len(source),
                        "user_content_lemma_types": len(user), "retained_types": len(source & user),
                        "source_type_retention": safe_rate(len(source & user), len(source), 1),
                        "jaccard": safe_rate(len(source & user), len(source | user), 1),
                        "tfidf_cosine_joint_40_documents": cosine,
                        "grounding_only_lemmas": ", ".join(sorted(source - user)),
                        "user_only_lemmas": ", ".join(sorted(user - source)),
                        "note": "Lexical overlap, NOT factual entailment or hallucination detection."})
    return frame, pd.DataFrame(lexical)


def repeated_phrases(sentences):
    # Normalized lexical n-grams: punctuation/placeholders omitted, never cross a sentence/turn.
    analyzer = CountVectorizer(analyzer="word", tokenizer=str.split, token_pattern=None,
                               lowercase=False, ngram_range=(3, 5)).build_analyzer()
    occurrences, case_sets = Counter(), defaultdict(set)
    for row in sentences:
        for phrase in analyzer(" ".join(row["tokens"])):
            key = (row["stage"], phrase)
            occurrences[key] += 1
            case_sets[key].add(row["seed_id"])
    rows = [{"stage": stage, "phrase": phrase, "distinct_cases": len(keys),
             "occurrences": occurrences[(stage, phrase)], "seed_ids": ", ".join(sorted(keys)),
             "note": "frequency only; repeated wording is not automatically a defect"}
            for (stage, phrase), keys in case_sets.items() if len(keys) >= 2]
    columns = ["stage", "phrase", "distinct_cases", "occurrences", "seed_ids", "note"]
    frame = pd.DataFrame(rows, columns=columns)
    return frame.sort_values(["distinct_cases", "occurrences", "phrase"], ascending=[False, False, True]) if rows else frame


def interaction_pairs(cases):
    rows = []
    for case in cases:
        for i, message in enumerate(case["messages"]):
            if message["role"] == "assistant" and "?" in message["content"]:
                reply = case["messages"][i + 1] if i + 1 < len(case["messages"]) else None
                adjacent = reply is not None and reply["role"] == "user"
                rows.append({"seed_id": case["seed_id"], "assistant_turn": i + 1,
                             "question_mark_present": True, "assistant_text": message["content"],
                             "following_user_turn_present": adjacent,
                             "following_user_text": reply["content"] if adjacent else "",
                             "adequacy": "NOT_ASSESSED", "note": "Question mark heuristic, not a speech-act classifier."})
    return pd.DataFrame(rows, columns=["seed_id", "assistant_turn", "question_mark_present", "assistant_text",
                                      "following_user_turn_present", "following_user_text", "adequacy", "note"])


def background_tables(root, features, contract):
    eda = Path(root) / EDA_REL
    historic = pd.read_csv(eda / "linguistic_summary.csv")
    historic = historic[historic.feature.isin(contract["patterns"])].copy()
    historic["instrument"] = "historical_v051_unchanged_regex"
    sensitivity = pd.read_csv(eda / "prevalence_sensitivity.csv")
    sensitivity = sensitivity[sensitivity.feature.isin(contract["patterns"])].copy()
    sensitivity["scope"] = "historical_population_only_separate_views_not_targets"
    rows = []
    for split in ["ALL_EXPLORATORY"] + sorted(features.release_split.unique()):
        selected = features if split == "ALL_EXPLORATORY" else features[features.release_split == split]
        for stage, group in selected.groupby("stage"):
            for name in contract["patterns"]:
                count = int(group["legacy_" + name].sum())
                rows.append({"release_split": split, "stage": stage, "feature": name,
                             "count": count, "n_cases": len(group), "rate": count / len(group),
                             "instrument": "historical_v051_unchanged_regex_on_unmodified_unit_text",
                             "scope": "descriptive_not_representative; any turn hit per case"})
    return historic, sensitivity, pd.DataFrame(rows)


def export_workbook(path, tables):
    wb = Workbook()
    wb.remove(wb.active)
    for name, frame in tables.items():
        ws = wb.create_sheet(name[:31])
        ws.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            values = [None if pd.isna(v) else v.item() if isinstance(v, np.generic) else v for v in row]
            if any(isinstance(v, str) and len(v) > 32767 for v in values):
                raise ValueError("Excel text would be truncated; inspect CSV/HTML instead")
            ws.append(values)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for i, column in enumerate(frame.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = 65 if any(x in column for x in ["text", "sentence", "note", "warning"]) else 27
            ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
            ws.cell(1, i).fill = PatternFill("solid", fgColor="1F4E78")
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Never execute source/generated text as a formula.
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def render_html(cases, summaries, lexical, phrases):
    parts = ["<!doctype html><html lang='zh-Hant'><meta charset='utf-8'><title>Language comparison</title>",
             "<style>body{font:16px system-ui;max-width:1550px;margin:auto;padding:2em}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit}"
             "table{border-collapse:collapse;display:block;overflow-x:auto}td,th{border:1px solid #ddd;padding:.5em}"
             ".cols{display:grid;grid-template-columns:repeat(3,1fr);gap:1em}section{background:#f5f6f8;padding:1em}"
             "@media(max-width:900px){.cols{display:block}}</style><h1>語言特徵比較 — 僅探索</h1>",
             "<ul>" + "".join("<li>" + html.escape(c) + "</li>" for c in CAUTIONS) + "</ul>",
             "<h2>角色分開的摘要（每案例等權）</h2>",
             summaries[(summaries.release_split == "ALL_EXPLORATORY") & summaries.feature.isin([
                 "word_count", "words_per_sentence", "negation_per100", "modal_tokens_per100", "mattr25_within_unit"])].to_html(index=False, escape=True),
             "<h2>Assistant 重複片語（前 20；不是缺陷判定）</h2>",
             phrases[phrases.stage == "generated_assistant"].head(20).to_html(index=False, escape=True)]
    lex_by_id = lexical.set_index("seed_id").to_dict("index")
    for case in cases:
        parts.append(f"<h2>{html.escape(case['seed_id'])}</h2><p>{html.escape(case['product'])} / {html.escape(case['issue'])}</p>")
        parts.append(f"<p>要求訊息數 {case['expected_messages']}；實際 {case['actual_messages']}。格式符合：{case['format_valid']}。詞彙重合不是事實正確率。</p>")
        parts.append("<div class='cols'>")
        for stage in ["grounding", "generated_user", "generated_assistant"]:
            parts.append(f"<section><h3>{stage}</h3><pre>" + html.escape("\n\n--- next turn ---\n\n".join(case["texts"][stage])) + "</pre></section>")
        parts.append("</div><p>User-only content lemmas（不是 hallucination 判定）: " + html.escape(lex_by_id[case["seed_id"]]["user_only_lemmas"]) + "</p>")
    return "\n".join(parts) + "</html>"


def run_analysis(project_root, out_dir, contract_path, run_rel=RUN_REL):
    root, out = Path(project_root).resolve(), Path(out_dir).resolve()
    allowed = root / "outputs/analysis/smoke_only/cfpb_v052_linguistic_comparison"
    if allowed not in out.parents:
        raise ValueError("Analysis output must be a new child of outputs/analysis/smoke_only/cfpb_v052_linguistic_comparison")
    if out.exists() and any(out.iterdir()):
        raise FileExistsError("Preserve the existing analysis; choose a new ANALYSIS_ID instead")
    inventory, raw, prepared = verify_inputs(root, run_rel)
    contract = read_json(contract_path)
    cases = assemble_cases(root, raw, prepared)
    nlp = load_parser()  # No regex-only substitute if the parser is unavailable.
    features, sentences, cues, lemmas, _ = extract_features(cases, nlp, contract)
    summaries = overview(features)
    differences, lexical = paired(features, lemmas)
    phrases = repeated_phrases(sentences)
    historic, sensitivity, current = background_tables(root, features, contract)
    case_info = pd.DataFrame([{k: v for k, v in c.items() if k not in {"texts", "messages"}} for c in cases])
    tables = {"Guide": pd.DataFrame({"warning": CAUTIONS}), "Case_inventory": case_info,
              "Side_by_side": pd.DataFrame([{"seed_id": c["seed_id"], "product": c["product"], "issue": c["issue"],
                  **{s: "\n\n--- next turn ---\n\n".join(c["texts"][s]) for s in ["grounding", "generated_user", "generated_assistant"]}}
                  for c in cases]),
              "Role_summary": summaries, "Case_features": features, "Paired_differences": differences,
              "Lexical_overlap": lexical, "Repeated_phrases": phrases,
              "Cue_evidence": cues, "Question_reply_pairs": interaction_pairs(cases),
              "EDA_legacy_rates": historic, "EDA_population_views": sensitivity, "Sample_legacy_rates": current,
              "EDA_TFIDF_context": pd.read_csv(root / EDA_REL / "tfidf_top_terms.csv"),
              "EDA_ngrams_context": pd.read_csv(root / EDA_REL / "frequent_ngrams.csv").head(100)}
    # Analyze first, create a NEW output directory only after input/parser validation.
    out.mkdir(parents=True, exist_ok=True)
    snapshot = out / "source_snapshot"
    snapshot.mkdir()
    (snapshot / Path(__file__).name).write_bytes(Path(__file__).read_bytes())
    (snapshot / "legacy_contract.json").write_bytes(Path(contract_path).read_bytes())
    for name, frame in tables.items():
        frame.to_csv(out / (name.lower() + ".csv"), index=False, encoding="utf-8-sig")
    export_workbook(out / "linguistic_comparison_20.xlsx", tables)
    (out / "linguistic_comparison_20.html").write_text(render_html(cases, summaries, lexical, phrases), encoding="utf-8")
    metadata = {"created_utc": datetime.now(timezone.utc).isoformat(), "source_run": run_rel,
                "config": CONFIG, "policy": POLICY, "input_hashes": inventory,
                "analysis_script_sha256": sha(__file__), "legacy_contract_sha256": sha(contract_path),
                "packages": {n: importlib.metadata.version(n) for n in ["spacy", "pandas", "pyarrow", "scikit-learn", "openpyxl"]},
                "cases": len(cases), "stage_rows": len(features), "all_cases_retained": True,
                "format_valid_cases": int(case_info.format_valid.sum()),
                "model_meta": {k: nlp.meta.get(k) for k in ["name", "version", "lang", "spacy_version"]}}
    # Verify inputs again after extraction/export; original runs are read-only.
    if any(sha(root / path) != expected for path, expected in inventory.items()):
        raise ValueError("An input changed during analysis; discard analysis interpretation")
    stats = summaries[summaries.release_split == "ALL_EXPLORATORY"]
    report = ["# 20 筆語言比較：描述性結果，不是品質判決", "",
              f"案例 {len(cases)}；保留全部原始案例。格式符合 {int(case_info.format_valid.sum())}/{len(cases)}。", "",
              "## 主要數值（每案例等權中位數）", ""]
    for feature in ["word_count", "words_per_sentence", "negation_per100", "modal_tokens_per100"]:
        values = stats[stats.feature == feature].set_index("stage")["median"]
        report.append("- " + feature + ": " + "; ".join(f"{s}={values[s]:.2f}" for s in STAGES))
    report += ["", "## 最常見的 assistant 片語（可重疊，不能將次數相加當缺陷數）", ""]
    for row in phrases[phrases.stage == "generated_assistant"].head(5).itertuples():
        report.append(f"- `{row.phrase}`：{row.distinct_cases}/20 個案例。")
    report += ["", "## 解讀順序", "", "- 預期變化：先看 full→excerpt→grounding 的長度／匿名化差異，不歸因於模型。",
               "- 待檢視變化：看 grounding→user 的配對差值與例句，判斷立場是否被過度肯定化；統計本身不能判定。",
               "- 模板化：看 assistant 跨案例片語，再看問句與下一輪；重複官方管道建議可能是 prompt 的預期效果。",
               "- 無法由此判定：法律／授權／事件真實性、隱私清除、實際客服互動自然度或母體代表性。", ""]
    report += ["- " + warning for warning in CAUTIONS]
    (out / "comparison_summary_zh.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    metadata["outputs"] = {p.relative_to(out).as_posix(): sha(p) for p in sorted(out.rglob("*")) if p.is_file()}
    save_json(out / "analysis_manifest.json", metadata)
    return metadata


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    parser.add_argument("--legacy-contract", type=Path, required=True)
    args = parser.parse_args()
    result = run_analysis(args.project_root, args.out_dir, args.legacy_contract)
    print(json.dumps({k: result[k] for k in ["cases", "stage_rows", "format_valid_cases", "policy"]}, indent=2))
