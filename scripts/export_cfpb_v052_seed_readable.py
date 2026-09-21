"""Export local, read-only Seed v05.2 reading copies; no API/network calls."""
from __future__ import annotations

import argparse
import hashlib
import html
import json
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "dataset/curated/seed_pools/cfpb_dispute/seed_v052"
LABELS = {"cfpb_seed_v052": "完整 Seed 主表", "cfpb_seed_v052_generation_input": "生成用摘錄與標籤",
          "cfpb_seed_v052_population_core": "比例抽樣核心", "cfpb_seed_v052_coverage_supplement": "分類覆蓋補充",
          "cfpb_seed_v052_enrichment": "Enrichment 補充", "cfpb_seed_v052_stress_test": "Stress test 壓力測試",
          "cfpb_seed_v052_excluded": "排除清單（原檔未提供完整敘事）"}
PRIORITY = ["seed_id", "record_id", "release_split", "Product", "product", "Sub-product", "sub_product",
            "Issue", "issue", "Sub-issue", "sub_issue", "seed_text", "seed_narrative_excerpt",
            "seed_action", "seed_action_reason", "enrichment_reason", "sampling_frame", "month"]
CAUTION = "私人閱讀副本；不是新 release。結構驗證不等於 privacy clearance；benchmark gate 仍關閉。"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def source_records(frame, iso=True):
    # pandas handles NumPy scalars, NaN/NaT, timestamps and nested arrays.
    return json.loads(frame.to_json(orient="records", date_format="iso" if iso else "epoch",
                                    date_unit="ms", force_ascii=False))


def verify_jsonl_pair(frame, jsonl):
    other = [json.loads(line) for line in jsonl.read_text(encoding="utf-8").splitlines() if line.strip()]
    if source_records(frame, iso=False) != other:
        raise ValueError(f"Parquet/JSONL differ; refuse to combine: {jsonl.name}")


def text_value(value):
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    return str(value)


def excel_text(value):
    return ILLEGAL_CHARACTERS_RE.sub(lambda m: f"\\u{ord(m.group()):04x}", value)


def export_excel(path, records, columns, title, sources):
    wb = Workbook(write_only=True)
    wrapping = Alignment(wrap_text=True, vertical="top")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    long_rows, references = [], {}
    controls = 0
    for number, record in enumerate(records, 1):
        for field, value in record.items():
            if not isinstance(value, str):
                continue
            cleaned = excel_text(value)
            controls += len(ILLEGAL_CHARACTERS_RE.findall(value))
            if len(cleaned) > 30000:
                ref = f"R{number}:{field}"
                references[(number, field)] = ref
                parts = [cleaned[i:i + 30000] for i in range(0, len(cleaned), 30000)]
                for i, part in enumerate(parts, 1):
                    long_rows.append([ref, number, record.get("seed_id", record.get("record_id", "")), field, i, len(parts), part])

    def sheet(name, headers, rows, widths=None):
        ws = wb.create_sheet(name)
        ws.freeze_panes = "B2"
        for i, field in enumerate(headers, 1):
            width = (widths or {}).get(field, 70 if any(t in field for t in ["text", "excerpt", "reason", "说明", "內容"]) else 25)
            ws.column_dimensions[get_column_letter(i)].width = width
        def append(values, head=False):
            cells = []
            for value in values:
                if isinstance(value, (dict, list)):
                    value = text_value(value)
                if isinstance(value, int) and not isinstance(value, bool) and abs(value) >= 10**15:
                    value = str(value)  # Excel loses precision on long identifiers.
                if isinstance(value, str):
                    value = excel_text(value)
                    if len(value) > 32767:
                        raise ValueError("Unexpected oversized Excel value; refuse truncation")
                cell = WriteOnlyCell(ws, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Source text can never become an Excel formula.
                cell.alignment = wrapping
                if head:
                    cell.font, cell.fill = header_font, header_fill
                cells.append(cell)
            ws.append(cells)
        append(headers, True)
        count = 1
        for values in rows:
            append(values)
            count += 1
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count}"

    def values(number, record, selected):
        return [f"完整長文見 Long_text：{references[(number, c)]}" if (number, c) in references
                else record.get(c) for c in selected]

    selected = [c for c in PRIORITY if c in columns]
    if not selected:
        selected = columns[:12]
    sheet("README", ["項目", "內容"], [
        ["資料", title], ["原始檔", ", ".join(sources)], ["列數", len(records)], ["欄位數", len(columns)],
        ["政策", CAUTION], ["Read_first", "常用識別、分類、敘事與篩選原因；未更動原始欄位名稱。"],
        ["All_fields", "全部原始欄位；timestamp 顯示 ISO UTC；null 為空白。需要型別精度請用原始 Parquet。"],
        ["Long_text", "超過 30,000 字元的欄位按 part 順序續接；沒有刪除長文。HTML 可直接讀全文。"],
        ["Column_catalog", "原始欄位名稱與非空數；正式意義仍以原 pipeline 文件為準。"],
        ["Excel 注意", "完整長文可能超過列高可顯示範圍；請用公式列或 HTML。公式樣式文字強制以純文字儲存。"],
        ["控制字元", f"Excel 不允許的控制字元以可見 Unicode escape 顯示：{controls} 處。HTML/原檔保留原文字。"],
        ["日期", "JSONL epoch-ms 與 Parquet timestamp 經內容核對後，閱讀版統一 ISO 日期。"],
        ["計數", "核心/覆蓋/enrichment 是主表子集；stress/excluded 是不同用途，勿將所有檔案列數相加當種子總數。"]])
    sheet("Read_first", ["source_row"] + selected,
          ([i] + values(i, r, selected) for i, r in enumerate(records, 1)))
    sheet("All_fields", ["source_row"] + columns,
          ([i] + values(i, r, columns) for i, r in enumerate(records, 1)))
    sheet("Long_text", ["reference", "source_row", "record_id", "field", "part", "parts", "text"], long_rows)
    sheet("Column_catalog", ["field", "non_null_rows", "display"],
          ([c, sum(r.get(c) is not None for r in records), "Original source field; see HTML for complete values"] for c in columns))
    counts = []
    for field in ["release_split", "Product", "product", "Issue", "issue"]:
        if field in columns:
            series = pd.Series([r.get(field) for r in records], dtype=object).fillna("(null)").value_counts(dropna=False)
            counts.extend([field, str(value), int(count)] for value, count in series.items())
    sheet("Counts", ["field", "value", "rows"], counts)
    wb.save(path)
    return {"long_text_fields": len(references), "long_text_parts": len(long_rows), "escaped_control_characters": controls}


BROWSER = r'''<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>__TITLE__</title>
<style>body{font:16px system-ui;max-width:1300px;margin:auto;padding:24px;color:#18212b;background:#f7f8fa}
header{position:sticky;top:0;background:#f7f8fa;padding:8px 0;z-index:2}input,select,button{font:inherit;padding:7px;margin:4px;max-width:95%}
input{width:440px}article{background:white;border:1px solid #ccd2da;border-radius:8px;padding:20px;margin:20px 0}
pre{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit;line-height:1.65}td,th{border:1px solid #ddd;padding:7px;text-align:left;vertical-align:top}
table{border-collapse:collapse;width:100%;table-layout:fixed}th{width:260px}small{color:#555}.warning{background:#fff4cf;padding:12px}
</style></head><body><p><a href="index.html">← 所有 Seed 閱讀版</a></p><h1>__TITLE__</h1>
<p class="warning">私人離線閱讀副本。不是 privacy clearance 或 benchmark release。篩選只改顯示，不修改原始資料。</p>
<header><input id="search" placeholder="搜尋 seed_id、標籤、敘事或其他欄位"><select id="split"><option value="">所有分組</option></select>
<select id="product"><option value="">所有 Product</option></select><select id="issue"><option value="">所有 Issue</option></select>
<div><button id="prev">上一頁</button><button id="next">下一頁</button><button id="reset">清除篩選</button><span id="status" aria-live="polite"></span></div></header>
<p>每頁 20 筆；敘事完整顯示，其他欄位可展開。「無敘事」表示此來源檔本來未提供，不是匯出遺失。</p><main id="results"></main>
<script id="data" type="application/json">__DATA__</script>
<script>
'use strict';
const rows=JSON.parse(document.getElementById('data').textContent), size=20;
const el=id=>document.getElementById(id), str=v=>v===null||v===undefined?'':typeof v==='object'?JSON.stringify(v):String(v);
const field=(r,a,b)=>str(r[a]??r[b]), searchable=rows.map(r=>Object.values(r).map(str).join('\n').toLowerCase());
let filtered=rows.map((_,i)=>i), page=0, timer;
function node(tag,text){const e=document.createElement(tag);if(text!==undefined)e.textContent=text;return e;}
function options(id,fn){[...new Set(rows.map(fn).filter(Boolean))].sort().forEach(v=>{const o=node('option',v);o.value=v;el(id).append(o);});}
options('split',r=>str(r.release_split));options('product',r=>field(r,'product','Product'));options('issue',r=>field(r,'issue','Issue'));
function render(){
  const pages=Math.max(1,Math.ceil(filtered.length/size));page=Math.max(0,Math.min(page,pages-1));
  el('status').textContent=` ${filtered.length} / ${rows.length} 筆；第 ${page+1} / ${pages} 頁`;
  el('prev').disabled=page===0;el('next').disabled=page>=pages-1;
  const fragment=document.createDocumentFragment();
  filtered.slice(page*size,(page+1)*size).forEach(i=>{const r=rows[i],article=node('article');
    article.append(node('h2',`${i+1}. ${str(r.seed_id??r.record_id??'(no ID)')}`));
    article.append(node('p',[str(r.release_split),field(r,'product','Product'),field(r,'issue','Issue')].filter(Boolean).join(' / ')));
    article.append(node('small',[field(r,'sub_product','Sub-product'),field(r,'sub_issue','Sub-issue')].filter(Boolean).join(' / ')));
    let hasText=false;
    ['seed_text','seed_narrative_excerpt'].forEach(k=>{if(Object.hasOwn(r,k)){hasText=true;article.append(node('h3',k),node('pre',str(r[k])));}});
    if(!hasText)article.append(node('p','此來源檔沒有 seed_text 或 seed_narrative_excerpt；下方列出所有原始欄位。'));
    const details=node('details'),table=node('table');details.append(node('summary','展開所有欄位（包含 EDA 特徵與篩選理由）'));
    Object.entries(r).forEach(([k,v])=>{const tr=node('tr'),td=node('td');td.append(node('pre',str(v)));tr.append(node('th',k),td);table.append(tr);});
    details.append(table);article.append(details);fragment.append(article);
  });
  el('results').replaceChildren(fragment);
}
function apply(){const query=el('search').value.trim().toLowerCase(),s=el('split').value,p=el('product').value,t=el('issue').value;
  filtered=rows.map((_,i)=>i).filter(i=>(!query||searchable[i].includes(query))&&(!s||str(rows[i].release_split)===s)&&(!p||field(rows[i],'product','Product')===p)&&(!t||field(rows[i],'issue','Issue')===t));page=0;render();}
el('search').addEventListener('input',()=>{clearTimeout(timer);timer=setTimeout(apply,180);});
['split','product','issue'].forEach(id=>el(id).addEventListener('change',apply));
el('prev').onclick=()=>{page--;render();};el('next').onclick=()=>{page++;render();};
el('reset').onclick=()=>{['search','split','product','issue'].forEach(id=>el(id).value='');apply();};render();
</script></body></html>'''


def render_browser(title, records):
    data = json.dumps(records, ensure_ascii=False, allow_nan=False, separators=(",", ":"))
    # Safe in a script-data element; source text is rendered only with textContent.
    data = data.replace("&", "\\u0026").replace("<", "\\u003c").replace(">", "\\u003e").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
    return BROWSER.replace("__TITLE__", html.escape(title)).replace("__DATA__", data)


def export_all(source=SOURCE, out=None):
    source = Path(source).resolve()
    out = Path(out).resolve() if out else source / "readable_v01"
    if source not in out.parents:
        raise ValueError("Use a new child directory of the source folder")
    if out.exists():
        raise FileExistsError("Reading copies already exist; use a new output directory")
    originals = {p.name: sha(p) for p in source.iterdir() if p.is_file()}
    parquet_files = sorted(source.glob("*.parquet"))
    if not parquet_files:
        raise ValueError("No Parquet files")
    unpaired = [p.name for p in source.glob("*.jsonl") if not p.with_suffix(".parquet").exists()]
    if unpaired:
        raise ValueError(f"Unpaired JSONL requires separate export: {unpaired}")
    loaded = []
    for path in parquet_files:
        frame = pd.read_parquet(path)
        sources = [path.name]
        if path.with_suffix(".jsonl").exists():
            verify_jsonl_pair(frame, path.with_suffix(".jsonl"))
            sources.append(path.with_suffix(".jsonl").name)
        loaded.append((path.stem, frame, sources))
    out.mkdir()
    datasets = []
    for name, frame, sources in loaded:
        records = source_records(frame)
        title = LABELS.get(name, name)
        info = export_excel(out / (name + "_readable.xlsx"), records, list(frame.columns), title, sources)
        (out / (name + "_browser.html")).write_text(render_browser(title, records), encoding="utf-8")
        entry = {"name": name, "title": title, "rows": len(frame), "columns": len(frame.columns),
                 "sources": sources, "equivalent_pair_verified": len(sources) == 2, **info}
        datasets.append(entry)
        print(json.dumps(entry, ensure_ascii=False), flush=True)
    links = []
    for d in datasets:
        name = d["name"]
        links.append(f'<tr><td>{html.escape(d["title"])}</td><td>{d["rows"]}</td><td>{d["columns"]}</td>'
                     f'<td><a href="{name}_browser.html">HTML 全文閱讀／搜尋</a></td>'
                     f'<td><a href="{name}_readable.xlsx">Excel 表格</a></td></tr>')
    (out / "index.html").write_text('<!doctype html><html lang="zh-Hant"><meta charset="utf-8"><title>Seed v05.2 閱讀入口</title>'
        '<style>body{font:17px system-ui;max-width:1100px;margin:40px auto}td,th{border:1px solid #ccc;padding:12px}table{border-collapse:collapse}</style>'
        '<h1>Seed v05.2 閱讀入口</h1><p>' + html.escape(CAUTION) + '</p>'
        '<p>主表 3,004 筆。子集與不同用途資料不可直接相加。所有轉換均離線，原始檔保留。</p>'
        '<p>用瀏覽器直接開啟 HTML，不需要 server。完整長文建議 HTML；Excel 的超長欄位在 Long_text 分頁續接。</p>'
        '<table><tr><th>資料</th><th>列數</th><th>欄位</th><th>閱讀</th><th>篩選</th></tr>' + ''.join(links) + '</table></html>', encoding="utf-8")
    readme = ["# CFPB Seed v05.2 閱讀副本", "", "先用瀏覽器開啟 index.html，選擇完整主表或生成用摘錄。", "", CAUTION, "",
              "- HTML：本機搜尋、分組／Product／Issue 篩選、每頁 20 筆，完整敘事及可展開原始欄位。無外部服務、CDN 或 API。",
              "- Excel：Read_first 常用欄位、All_fields 全欄位、Counts 分布、Column_catalog 欄位清單、Long_text 長文續頁。",
              "- 凍結首列與首欄、自動篩選、文字換行。Excel 列高可能無法顯示全文，請改用 HTML。",
              "- 大於 30,000 字元的欄位在 Long_text 按 reference / part 續接；不默默截斷。",
              "- 日期顯示 ISO 格式；null 為空白；Excel 不允許的控制字元顯示為 Unicode escape。",
              "- 同名 Parquet/JSONL 已核對每列每欄一致（timestamp 以 epoch-ms 對齊），故共用一份閱讀版。",
              "- 閱讀版不是機器可逆的新資料 release；統計或精度敏感工作仍用原始 Parquet／JSONL。",
              "- excluded 原始檔未提供完整敘事，匯出不從其他來源回填、不捏造。",
              "- HTML/Excel 仍含文本與衍生資料，只在私人位置使用，不上傳線上轉檔網站或公開分享。", "",
              "## 檔案對照", ""]
    readme += [f'- {", ".join(d["sources"])} → {d["name"]}_readable.xlsx / {d["name"]}_browser.html（{d["rows"]} 筆）' for d in datasets]
    (out / "README_zh.md").write_text("\n".join(readme) + "\n", encoding="utf-8")
    if {p.name: sha(p) for p in source.iterdir() if p.is_file()} != originals:
        raise ValueError("Source files changed during export")
    manifest = {"created_utc": datetime.now(timezone.utc).isoformat(), "source_directory": source.name,
                "original_files_sha256": originals, "source_files_unchanged": True, "datasets": datasets,
                "script_sha256": sha(__file__), "privacy_clearance_claimed": False,
                "outputs": {p.name: sha(p) for p in out.iterdir() if p.is_file()}}
    write_json(out / "export_manifest.json", manifest)
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    export_all(out=args.out)
