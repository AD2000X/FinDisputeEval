"""Build a review-ready Excel workbook for the 20-row v02 adjudication set."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


_PROJECT_ROOT_ENV = os.environ.get("FINDISPUTEEVAL_PROJECT_ROOT")
PROJECT_ROOT = (
    Path(_PROJECT_ROOT_ENV).resolve()
    if _PROJECT_ROOT_ENV
    else Path(__file__).resolve().parents[3]
)
SMOKE_ROOT = PROJECT_ROOT / "outputs/generation/smoke_only/cfpb_seed_v052_pipeline_override"
RUN_ROOT = SMOKE_ROOT / "run_20260722T135306Z"
REVALIDATION_ROOT = (
    SMOKE_ROOT / "revalidations/validator_v02/source_run_20260722T135306Z"
)
DEFAULT_TEMPLATE = REVALIDATION_ROOT / "oracle/human_adjudication_20_template.csv"
DEFAULT_PREPARED = RUN_ROOT / "prepared_inputs/nemo_seed_v052_pipeline_smoke_20.jsonl"
DEFAULT_OUTPUT = REVALIDATION_ROOT / "oracle/human_adjudication_20_review.xlsx"

REVIEW_COLUMNS = [
    "seed_id",
    "product",
    "sub_product",
    "issue",
    "sub_issue",
    "grounding_evidence",
    "generated_evidence",
    "candidate_v02_reasons_json",
    "expected_decision",
    "expected_v02_reasons_json",
    "reviewer_id",
    "reviewed_utc",
    "review_status",
]

COLUMN_WIDTHS = {
    "seed_id": 30,
    "product": 30,
    "sub_product": 38,
    "issue": 42,
    "sub_issue": 44,
    "grounding_evidence": 72,
    "generated_evidence": 90,
    "candidate_v02_reasons_json": 46,
    "expected_decision": 20,
    "expected_v02_reasons_json": 48,
    "reviewer_id": 20,
    "reviewed_utc": 28,
    "review_status": 18,
}

REASON_CODES = [
    ("insufficient_grounding", "Post-redaction facts are too sparse for safe generation."),
    ("unsupported_scenario_detail", "Material generated fact is absent from grounding."),
    ("authorization_state_changed", "Authorization status changed or was asserted from ambiguity."),
    ("claim_type_changed", "The generated dispute type materially differs from grounding."),
    ("factual_status_changed", "Paid/resolved/pending/accurate status materially changed."),
    ("label_grounding_conflict", "Classification metadata conflicts with narrative grounding."),
    ("legal_or_rights_claim", "Unsupported legal right or conclusion is asserted."),
    ("unsupported_product_policy", "Product rule, fee cause, or closure cause is invented."),
    ("unsupported_procedural_guidance", "Institution-specific process or remedy is invented."),
    ("indirect_sensitive_information_guidance", "Guidance encourages sensitive account information."),
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_review_frame(template_path: Path, prepared_path: Path) -> pd.DataFrame:
    template = pd.read_csv(
        template_path,
        dtype=str,
        keep_default_na=False,
        encoding="utf-8-sig",
    )
    prepared = pd.read_json(prepared_path, lines=True, dtype=False)
    if template["seed_id"].duplicated().any() or prepared["seed_id"].duplicated().any():
        raise ValueError("Template or prepared input has duplicate seed_id values")
    if set(template["seed_id"].astype(str)) != set(prepared["seed_id"].astype(str)):
        raise ValueError("Template and prepared input seed coverage differ")
    labels = prepared[
        ["seed_id", "product", "sub_product", "issue", "sub_issue"]
    ].copy()
    for column in labels:
        labels[column] = labels[column].fillna("").astype(str)
    merged = template.merge(labels, on="seed_id", how="left", validate="one_to_one")
    missing = sorted(set(REVIEW_COLUMNS) - set(merged.columns))
    if missing:
        raise ValueError(f"Review workbook columns are missing: {missing}")
    return merged[REVIEW_COLUMNS].copy()


def add_review_sheet(workbook: Workbook, frame: pd.DataFrame) -> None:
    sheet = workbook.active
    sheet.title = "Review"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    reference_fill = PatternFill("solid", fgColor="E7E6E6")
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    reviewed_fill = PatternFill("solid", fgColor="E2F0D9")
    thin_gray = Side(style="thin", color="D9E1F2")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for column_index, column in enumerate(REVIEW_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=column)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        sheet.column_dimensions[get_column_letter(column_index)].width = COLUMN_WIDTHS[column]
    sheet.row_dimensions[1].height = 34

    editable_columns = {
        "expected_decision", "expected_v02_reasons_json", "reviewer_id",
        "reviewed_utc", "review_status",
    }
    for row_index, record in enumerate(frame.to_dict(orient="records"), start=2):
        for column_index, column in enumerate(REVIEW_COLUMNS, start=1):
            value: Any = record[column]
            if pd.isna(value):
                value = ""
            cell = sheet.cell(row=row_index, column=column_index, value=str(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            cell.fill = editable_fill if column in editable_columns else reference_fill
        sheet.row_dimensions[row_index].height = 126

    last_row = len(frame) + 1
    last_column = get_column_letter(len(REVIEW_COLUMNS))
    sheet.auto_filter.ref = f"A1:{last_column}{last_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    decision_column = get_column_letter(REVIEW_COLUMNS.index("expected_decision") + 1)
    status_column = get_column_letter(REVIEW_COLUMNS.index("review_status") + 1)
    decision_validation = DataValidation(
        type="list", formula1='"accept,reject,review"', allow_blank=True
    )
    status_validation = DataValidation(
        type="list", formula1='"pending,reviewed"', allow_blank=False
    )
    decision_validation.error = "Select accept, reject, or review."
    decision_validation.errorTitle = "Invalid decision"
    decision_validation.prompt = "For the completed oracle, resolve rows to accept or reject."
    decision_validation.promptTitle = "Expected decision"
    decision_validation.showErrorMessage = True
    decision_validation.showInputMessage = True
    status_validation.showErrorMessage = True
    sheet.add_data_validation(decision_validation)
    sheet.add_data_validation(status_validation)
    decision_validation.add(f"{decision_column}2:{decision_column}{last_row}")
    status_validation.add(f"{status_column}2:{status_column}{last_row}")

    sheet.conditional_formatting.add(
        f"A2:{last_column}{last_row}",
        FormulaRule(
            formula=[f'${status_column}2="reviewed"'],
            fill=reviewed_fill,
        ),
    )


def add_instructions_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Instructions")
    sheet.sheet_view.showGridLines = False
    instructions = [
        ("Purpose", "Human ground-truth adjudication; candidate reasons are analysis leads only."),
        ("Evidence priority", "Judge only the grounding and labels visible to the model, then generated content."),
        ("Accept", "Use expected_decision=accept and expected_v02_reasons_json=[]."),
        ("Reject", "Use expected_decision=reject and the smallest complete JSON list of supported reasons."),
        ("Completion", "Set reviewer_id, ISO-8601 reviewed_utc, and review_status=reviewed."),
        ("Important", "Do not use the same human answers as independent judge predictions."),
    ]
    sheet.append(["Item", "Instruction"])
    for item in instructions:
        sheet.append(item)
    sheet.append([])
    sheet.append(["Reason code", "Definition"])
    for reason in REASON_CODES:
        sheet.append(reason)
    for cell in sheet[1] + sheet[9]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 100
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"


def add_lineage_sheet(
    workbook: Workbook,
    template_path: Path,
    prepared_path: Path,
) -> None:
    sheet = workbook.create_sheet("_lineage")
    rows = [
        ("artifact", "path", "sha256"),
        ("oracle_template", str(template_path.resolve()), sha256_file(template_path)),
        ("prepared_input", str(prepared_path.resolve()), sha256_file(prepared_path)),
    ]
    for row in rows:
        sheet.append(row)
    sheet.sheet_state = "hidden"


def build_workbook(template_path: Path, prepared_path: Path, output_path: Path) -> dict[str, Any]:
    input_hashes_before = {
        "template": sha256_file(template_path),
        "prepared": sha256_file(prepared_path),
    }
    frame = load_review_frame(template_path, prepared_path)
    if len(frame) != 20:
        raise ValueError(f"Expected 20 review rows, received {len(frame)}")

    workbook = Workbook()
    add_review_sheet(workbook, frame)
    add_instructions_sheet(workbook)
    add_lineage_sheet(workbook, template_path, prepared_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    input_hashes_after = {
        "template": sha256_file(template_path),
        "prepared": sha256_file(prepared_path),
    }
    if input_hashes_after != input_hashes_before:
        raise RuntimeError("An input artifact changed while the workbook was built")

    loaded = load_workbook(output_path, read_only=False, data_only=False)
    review = loaded["Review"]
    if review.freeze_panes != "A2":
        raise RuntimeError("Workbook verification failed: first row is not frozen")
    if review.auto_filter.ref != f"A1:M{len(frame) + 1}":
        raise RuntimeError("Workbook verification failed: auto-filter range")
    if review.max_row != 21 or review.max_column != len(REVIEW_COLUMNS):
        raise RuntimeError("Workbook verification failed: shape")
    if len(review.data_validations.dataValidation) != 2:
        raise RuntimeError("Workbook verification failed: dropdown validations")
    headers = [cell.value for cell in review[1]]
    if headers != REVIEW_COLUMNS:
        raise RuntimeError("Workbook verification failed: column order")

    return {
        "output": str(output_path.resolve()),
        "rows": len(frame),
        "columns": REVIEW_COLUMNS,
        "sheets": loaded.sheetnames,
        "freeze_panes": review.freeze_panes,
        "auto_filter": review.auto_filter.ref,
        "data_validations": len(review.data_validations.dataValidation),
        "sha256": sha256_file(output_path),
        "size_bytes": output_path.stat().st_size,
        "inputs_unchanged": True,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--prepared", type=Path, default=DEFAULT_PREPARED)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    result = build_workbook(args.template, args.prepared, args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
