"""Offline preparation, immutable run bookkeeping, and review-first exports.

No judge or provider calls live in this module. Existing v02 privacy and
deterministic checks are reused; a clean deterministic check is NOT acceptance.
"""
from __future__ import annotations

import html
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

try:
    from . import prepare_cfpb_seed_v052_pipeline_smoke_v02 as prep
    from .cfpb_v052_pipeline_smoke_v02_common import load_records, sha256_file
    from .validate_cfpb_v052_pipeline_smoke_v02 import parse_structured, validate_row
except ImportError:
    import prepare_cfpb_seed_v052_pipeline_smoke_v02 as prep
    from cfpb_v052_pipeline_smoke_v02_common import load_records, sha256_file
    from validate_cfpb_v052_pipeline_smoke_v02 import parse_structured, validate_row

POLICY = dict(pipeline_smoke_only=True, provisional=True, privacy_verified=False,
              benchmark_eligible=False, publication_eligible=False,
              may_enter_final_dataset=False, distribution_calibration_eligible=False,
              formal_pilot_allowed=False, judge_frozen=False)
REFERENCE_COLUMNS = ["seed_id", "product", "sub_product", "issue", "sub_issue",
                     "grounding_evidence", "generated_dialogue", "generated_summary",
                     "generated_privacy_notes", "deterministic_flags_json"]
DIMENSIONS = ["factual_fidelity", "unsupported_content", "dialogue_quality", "task_usefulness"]
EDIT_COLUMNS = DIMENSIONS + ["review_decision", "problem_evidence", "revision_suggestion",
                             "reviewer_kind", "reviewer_id", "reviewed_utc", "review_status"]
COLUMNS = REFERENCE_COLUMNS + EDIT_COLUMNS
CHOICES = {**{key: ["pass", "fail", "uncertain"] for key in DIMENSIONS},
           "review_decision": ["usable", "needs_revision", "unusable", "uncertain"],
           "reviewer_kind": ["human", "llm"], "review_status": ["pending", "reviewed"]}


def utc_now():
    return datetime.now(timezone.utc).isoformat()


def write_once(path: Path, content: str | bytes):
    """Allow identical resumes, never replace an existing different artifact."""
    payload = content.encode("utf-8") if isinstance(content, str) else content
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        if path.read_bytes() != payload:
            raise ValueError(f"Existing artifact differs; use a new run: {path}")
        return
    with path.open("xb") as handle:
        handle.write(payload)


def save_json(path: Path, value):
    write_once(path, json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n")


def verify_inventory(project_root: Path, inventory: dict):
    for relative, expected in inventory.items():
        path = project_root / relative
        if not path.is_file():
            raise FileNotFoundError(f"Sync required input: {relative}")
        if sha256_file(path) != expected:
            raise ValueError(f"Input SHA-256 mismatch: {relative}")


def check_coverage(raw_path: Path, prepared_path: Path):
    seeds, raw = load_records(prepared_path), load_records(raw_path)
    for label, rows in (("prepared", seeds), ("raw", raw)):
        ids = [row.get("seed_id") for row in rows]
        if any(not isinstance(key, str) or not key for key in ids) or len(set(ids)) != len(ids):
            raise ValueError(f"{label} has missing or duplicate seed IDs")
    if {r["seed_id"] for r in raw} != {r["seed_id"] for r in seeds}:
        raise ValueError("Raw/prepared coverage mismatch; partial results preserved, no automatic regeneration")
    return seeds, raw


def prepare_run(project_root: Path, run_root: Path, input_inventory: dict, random_seed=20260919):
    verify_inventory(project_root, input_inventory)
    seed_root = project_root / "dataset/curated/seed_pools/cfpb_dispute/seed_v052"
    privacy = project_root / ("dataset/curated/annotations/cfpb_seed_v05_audit/"
                              "run_20260713T145423Z/privacy_qa/full_v052_v02/")
    old = project_root / ("outputs/generation/smoke_only/cfpb_seed_v052_pipeline_override/"
                          "run_20260722T135306Z/prepared_inputs/nemo_seed_v052_pipeline_smoke_20.jsonl")
    old_rows = load_records(old)
    if len(old_rows) != 20 or len({r["seed_id"] for r in old_rows}) != 20:
        raise ValueError("Expected exactly 20 unique prior development seeds")
    output = run_root / "prepared_inputs"
    sample = output / "nemo_seed_v052_pipeline_smoke_20.jsonl"
    manifest = output / "pipeline_smoke_input_manifest.json"
    if manifest.exists():
        info = json.loads(manifest.read_text(encoding="utf-8"))
        if (info["random_seed"] != random_seed or info["output"]["sha256"] != sha256_file(sample)
                or info["prior_sample_exclusion"]["source_sha256"] != sha256_file(old)):
            raise ValueError("Prepared cache differs from this run")
    else:
        if output.exists() and any(output.iterdir()):
            raise ValueError("Partial prepared input exists; preserve it and choose a new run")
        sample, manifest, info = prep.prepare_smoke_inputs(
            seed_input=seed_root / "cfpb_seed_v052_generation_input.jsonl",
            full_seed_path=seed_root / "cfpb_seed_v052.parquet",
            seed_manifest=seed_root / "seed_v052_manifest.json",
            disposition_path=privacy / "pipeline_privacy_disposition_v01.json",
            output_dir=output, sample_size=20, random_seed=random_seed,
            exclude_prepared_input=old)
    seeds = load_records(sample)
    if len(seeds) != 20 or len({r["seed_id"] for r in seeds}) != 20:
        raise ValueError("Expected exactly 20 unique prepared rows")
    if {r["seed_id"] for r in seeds} & {r["seed_id"] for r in old_rows}:
        raise ValueError("Prior development seed overlap")
    for row in seeds:
        if (row.get("offset_invariant_verified") is not True
                or row.get("privacy_verified") is not False
                or row.get("benchmark_eligible") is not False
                or "seed_text" in row or "seed_narrative_excerpt" in row):
            raise ValueError("Unsafe prepared input contract")
    return sample, manifest, info


def resolve_cached_raw(run_root: Path, prepared: Path):
    """Never regenerate implicitly after any API attempt, even if it failed."""
    pointer_path = run_root / "raw_output_pointer.json"
    if pointer_path.exists():
        pointer = json.loads(pointer_path.read_text(encoding="utf-8"))
        raw = (run_root / pointer["path"]).resolve()
        if run_root.resolve() not in raw.parents or sha256_file(raw) != pointer["sha256"]:
            raise ValueError("Raw cache path/hash mismatch")
        if sha256_file(prepared) != pointer["prepared_sha256"]:
            raise ValueError("Prepared input differs from raw cache provenance")
        check_coverage(raw, prepared)
        return raw
    candidates = sorted((run_root / "raw/data_designer/dataset/parquet-files").glob("*.parquet"))
    if candidates:
        if len(candidates) != 1:
            raise ValueError("Ambiguous raw batches; inspect without rerunning API")
        check_coverage(candidates[0], prepared)
        record_raw(run_root, candidates[0], prepared)
        return candidates[0]
    if (run_root / "generation_started.json").exists():
        raise RuntimeError("Previous API attempt has no complete batch. Preserve this run; inspect before a new run.")
    raw_root = run_root / "raw/data_designer"
    if raw_root.exists() and any(raw_root.iterdir()):
        raise RuntimeError("Untracked raw artifacts exist; refusing another API call")
    return None


def record_raw(run_root: Path, raw: Path, prepared: Path):
    check_coverage(raw, prepared)
    relative = raw.resolve().relative_to(run_root.resolve()).as_posix()
    save_json(run_root / "raw_output_pointer.json", {
        "path": relative, "sha256": sha256_file(raw), "prepared_sha256": sha256_file(prepared)})


def review_records(raw_path: Path, prepared_path: Path):
    seeds, raw = check_coverage(raw_path, prepared_path)
    raw_by_id = {row["seed_id"]: row for row in raw}
    seed_by_id = {row["seed_id"]: row for row in seeds}
    records = []
    for seed in seeds:
        row = raw_by_id[seed["seed_id"]]
        reasons, _, _ = validate_row(row, seed_by_id)  # No semantic judgment supplied.
        obj = parse_structured(row.get("dialogue"))
        messages = obj.get("conversation")
        if isinstance(messages, list) and all(isinstance(m, dict) for m in messages):
            dialogue = "\n\n".join(f"{i}. {str(m.get('role', '?')).upper()}\n{m.get('content', '')}"
                                      for i, m in enumerate(messages, 1))
        else:
            dialogue = json.dumps(row.get("dialogue"), ensure_ascii=False, indent=2)
        record = {key: str(seed.get(key, "")) for key in REFERENCE_COLUMNS[:5]}
        record.update(grounding_evidence=seed["generation_grounding_excerpt"],
                      generated_dialogue=dialogue,
                      generated_summary=str(obj.get("synthetic_case_summary", "")),
                      generated_privacy_notes=json.dumps(obj.get("privacy_notes", []), ensure_ascii=False),
                      deterministic_flags_json=json.dumps(reasons),
                      **{key: "" for key in EDIT_COLUMNS})
        record["review_status"] = "pending"
        records.append(record)
    return records


def workbook_values(path: Path):
    wb = load_workbook(path, data_only=False)
    sheet = wb["Review"]
    rows = list(sheet.values)
    if list(rows[0]) != COLUMNS:
        raise ValueError("Review workbook columns changed")
    return [dict(zip(COLUMNS, ("" if v is None else str(v) for v in row))) for row in rows[1:]]


def verify_reference_rows(actual, expected):
    if len(actual) != len(expected) or len({r["seed_id"] for r in actual}) != len(expected):
        raise ValueError("Review workbook row coverage changed")
    by_id = {row["seed_id"]: row for row in actual}
    for ref in expected:
        row = by_id.get(ref["seed_id"], {})
        if any(row.get(key) != ref[key] for key in REFERENCE_COLUMNS):
            raise ValueError(f"Reference evidence changed: {ref['seed_id']}")


def make_workbook(path: Path, records, lineage):
    if path.exists():
        verify_reference_rows(workbook_values(path), records)
        return  # Preserve all user annotations.
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.append(COLUMNS)
    for record in records:
        if any(len(str(record[key])) > 32767 for key in COLUMNS):
            raise ValueError("Excel cell limit exceeded; refusing silent truncation")
        ws.append([record[key] for key in COLUMNS])
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
    for col, key in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col)].width = (
            70 if key in {"grounding_evidence", "generated_dialogue"} else
            42 if key in {"generated_summary", "problem_evidence", "revision_suggestion"} else 24)
        for row in ws.iter_rows(min_col=col, max_col=col):
            cell = row[0]
            if isinstance(cell.value, str):
                if len(cell.value) > 32767:
                    raise ValueError("Excel cell limit exceeded; use raw/HTML, never silently truncate")
                cell.data_type = "s"  # Untrusted text is never an Excel formula.
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = PatternFill("solid", fgColor="FFF2CC" if key in EDIT_COLUMNS else "E7E6E6")
        ws.cell(1, col).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, col).fill = PatternFill("solid", fgColor="1F4E78")
        if key in CHOICES:
            dv = DataValidation(type="list", formula1='"' + ','.join(CHOICES[key]) + '"', allow_blank=True)
            dv.showErrorMessage = True
            dv.errorStyle = "stop"
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{len(records)+1}")
    for i in range(2, len(records) + 2):
        ws.row_dimensions[i].height = 240
    guide = wb.create_sheet("Guide")
    for line in [
        "探索性 smoke；不是 oracle、人工 gold、privacy clearance 或 benchmark。",
        "灰色欄位是來源證據，不可修改；黃色欄位由 reviewer 填寫。",
        "先讀 grounding，再看 dialogue。分類是 metadata，不是絕對真相。",
        "factual_fidelity: 是否保持授權狀態、爭議類型、事件與付款狀態？",
        "unsupported_content: pass = 沒有發現重大虛構政策、權利、承諾或細節。",
        "dialogue_quality: 是否自然、有實質互動，而非重複客服模板？",
        "task_usefulness: 是否適合作為金融爭議 intake 的探索性合成對話？",
        "usable=原樣可用於探索；needs_revision=可修；unusable=需重做；uncertain=需討論。",
        "usable 只表示本次品質閱讀判斷；不會解除 smoke-only 或 privacy gate。",
        "所有四面向都填 pass/fail/uncertain；非 usable 請填問題片段/說明。",
        "reviewer_kind 必須標 human 或 llm；AI 初評不能冒充 human。",
        "填 reviewer_id 與 UTC ISO 時間，例如 2026-09-19T12:00:00Z，再設 reviewed。",
        "deterministic_flags 僅是既有規則的提醒；空白問題集合不代表語意合格。",
        "長文字可能超出 Excel 列高；請用 review_readable.html 看完整內容。",
        "儲存並同步 Excel 後，重跑最後的摘要 cell；它不會呼叫模型或改寫 Excel。",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 115
    for row in guide:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    meta = wb.create_sheet("_lineage")
    meta.append([json.dumps(lineage, sort_keys=True)])
    meta.sheet_state = "hidden"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_review_pack(raw_path: Path, prepared_path: Path, out_dir: Path):
    records = review_records(raw_path, prepared_path)
    lineage = {"raw_sha256": sha256_file(raw_path), "prepared_sha256": sha256_file(prepared_path),
               "policy": POLICY, "semantic_judge_used": False}
    save_json(out_dir / "review_reference.json", {"lineage": lineage, "rows": records})
    workbook = out_dir / "generation_review_20.xlsx"
    make_workbook(workbook, records, lineage)
    parts = ["<!doctype html><html lang='zh-Hant'><meta charset='utf-8'><title>Generation review</title>",
             "<style>body{font:16px system-ui;max-width:1500px;margin:2em auto;padding:1em}"
             ".pair{display:grid;grid-template-columns:1fr 1fr;gap:2em}pre{white-space:pre-wrap;"
             "overflow-wrap:anywhere;font:inherit;background:#f3f4f6;padding:1em}"
             "article{border-top:2px solid #bbb;margin-top:2em}@media(max-width:800px){.pair{display:block}}</style>",
             "<h1>20-row generation review — smoke only</h1><p>Privacy unverified. No LLM judge. "
             "All rows pending review. Enter decisions in Excel. Source allegations are not verified facts.</p>"]
    for row in records:
        esc = lambda key: html.escape(row[key])
        parts.append(f"<article><h2>{esc('seed_id')}</h2><p>" + " / ".join(esc(k) for k in REFERENCE_COLUMNS[1:5])
                     + f"</p><div class='pair'><section><h3>Grounding</h3><pre>{esc('grounding_evidence')}</pre>"
                     + f"</section><section><h3>Generated dialogue</h3><pre>{esc('generated_dialogue')}</pre>"
                     + f"<h3>Generated summary</h3><pre>{esc('generated_summary')}</pre></section></div>"
                     + f"<p>Deterministic flags: {esc('deterministic_flags_json')}</p></article>")
    parts.append("</html>")
    write_once(out_dir / "review_readable.html", "\n".join(parts))
    counts = Counter(reason for row in records for reason in json.loads(row["deterministic_flags_json"]))
    save_json(out_dir / "technical_check_report.json", {
        **lineage, "raw_rows": len(records), "pending_review_rows": len(records),
        "rows_with_deterministic_flags": sum(bool(json.loads(r["deterministic_flags_json"])) for r in records),
        "deterministic_reason_counts": dict(counts), "auto_accepted_rows": 0,
        "note": "All rows retained for reading, including malformed or flagged outputs."})
    return workbook


def summarize_review(workbook: Path):
    reference = json.loads((workbook.parent / "review_reference.json").read_text(encoding="utf-8"))
    records = workbook_values(workbook)
    verify_reference_rows(records, reference["rows"])
    reviewed = []
    for row in records:
        for key, choices in CHOICES.items():
            if row[key] and row[key] not in choices:
                raise ValueError(f"Invalid {key}: {row['seed_id']}")
        if row["review_status"] != "reviewed":
            continue
        required = DIMENSIONS + ["review_decision", "reviewer_kind", "reviewer_id", "reviewed_utc"]
        if any(not row[key].strip() for key in required):
            raise ValueError(f"Incomplete review provenance/dimensions: {row['seed_id']}")
        when = datetime.fromisoformat(row["reviewed_utc"].replace("Z", "+00:00"))
        if when.utcoffset() is None or when.utcoffset().total_seconds() != 0:
            raise ValueError("reviewed_utc must be an explicit UTC timestamp")
        if row["review_decision"] == "usable" and any(row[k] != "pass" for k in DIMENSIONS):
            raise ValueError("usable requires all four dimensions to pass")
        if row["review_decision"] != "usable" and not row["problem_evidence"].strip():
            raise ValueError("Non-usable review requires a problem excerpt or uncertainty explanation")
        reviewed.append(row)
    report = {"workbook_sha256": sha256_file(workbook), "rows": len(records),
              "reviewed_rows": len(reviewed), "pending_rows": len(records) - len(reviewed),
              "decisions_by_reviewer_kind": {
                  kind: dict(Counter(r["review_decision"] for r in reviewed if r["reviewer_kind"] == kind))
                  for kind in ["human", "llm"]},
              "failed_dimensions": dict(Counter(k for r in reviewed for k in DIMENSIONS if r[k] == "fail")),
              "revision_suggestions": [r["revision_suggestion"] for r in reviewed if r["revision_suggestion"]],
              "policy": POLICY, "note": "Exploration only; no judge metrics, no population estimate, no automatic promotion."}
    stem = workbook.parent / ("summary_" + report["workbook_sha256"][:12])
    save_json(stem.with_suffix(".json"), report)
    summary = ["# Nemotron 探索性生成摘要", "", f"已審 {len(reviewed)}/{len(records)}；待審 {report['pending_rows']}。",
               "", "本摘要只統計已完成的 reviewer 紀錄；pending 不算可用。AI 與人工分開統計。", ""]
    for kind, counts in report["decisions_by_reviewer_kind"].items():
        summary.append(f"- {kind}: " + (json.dumps(counts, ensure_ascii=False) if counts else "尚無完成的評閱"))
    summary.extend(["", "## 最常出現的失敗面向（不是完整缺陷 taxonomy）", "",
                    json.dumps(report["failed_dimensions"], ensure_ascii=False), "", "## Reviewer 修改建議", ""])
    summary.extend([f"- {s}" for s in report["revision_suggestions"]] or ["尚無；請先閱讀 Excel/HTML。"])
    summary.extend(["", "下一步：讀完這 20 筆後，只選最重要的 1–2 項修改生成方式，不自動擴大筆數。",
                    "privacy_verified=false；benchmark_eligible=false；非 formal pilot；judge 校準暫停。"])
    write_once(stem.with_suffix(".md"), "\n".join(summary) + "\n")
    return report
