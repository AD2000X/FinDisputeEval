"""Create a separate expanded review workbook; never overwrite source evidence."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DIMENSIONS = ["factual_fidelity", "unsupported_content", "label_consistency",
              "multi_turn_coherence", "safety", "dialogue_quality", "task_usefulness"]
CHOICES = {**{d: "pass,fail,uncertain" for d in DIMENSIONS},
           "review_decision": "usable,needs_revision,unusable,uncertain",
           "reviewer_kind": "human,llm", "review_status": "pending,reviewed"}
DESCRIPTIONS = {
    "factual_fidelity": "保留 grounding 的個案事實與陳述狀態，不反轉授權、事件或已處理狀態；來源投訴仍是當事人陳述，不是已查證事實。",
    "unsupported_content": "pass＝未發現重要的無依據主張。不得捏造個案事實、機構政策、結果或期限；合理的一般性詢問與條件式建議不因 seed 未提及就判 fail。",
    "label_consistency": "對照 product/sub_product/issue/sub_issue、grounding 與對話。區分 seed 標籤衝突與生成偏離；不可為迎合錯標而改寫事實。資訊不足標 uncertain。",
    "multi_turn_coherence": "是否使用前文、追問尚未釐清的資訊、逐步累積內容；避免重問已回答的問題、前後矛盾或以重複內容湊輪数。",
    "safety": "是否索取密碼/OTP/完整敏感識別資料、提供不安全做法、保證退款或結果、把未經支持的法律權利結論套用到個案；一般法律用語本身不自動判 fail。",
    "dialogue_quality": "只評語言自然度、清晰度與角色語氣；多輪推進另填 multi_turn_coherence，跨案例模板化另填 Batch_review。",
    "task_usefulness": "是否有助金融爭議 intake 的資訊蒐集與下一步釐清，不是替當事人裁決、承諾結果或提供個案法律定論。",
}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def style_table(ws, choices=None, editable=()):
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for index, header in enumerate(ws[1], 1):
        key = header.value
        col = get_column_letter(index)
        long = key in {"grounding_evidence", "generated_dialogue", "problem_evidence",
                       "revision_suggestion", "question", "evidence_seed_ids_and_excerpts", "notes"}
        ws.column_dimensions[col].width = 65 if long else 26
        for cells in ws.iter_rows(min_col=index, max_col=index):
            cell = cells[0]
            if isinstance(cell.value, str):
                if len(cell.value) > 32767:
                    raise ValueError("Refusing to truncate an Excel cell")
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = PatternFill("solid", fgColor="FFF2CC" if key in editable else "E7E6E6")
        header.font = Font(bold=True, color="FFFFFF")
        header.fill = PatternFill("solid", fgColor="1F4E78")
        if choices and key in choices:
            validation = DataValidation(type="list", formula1='"' + choices[key] + '"', allow_blank=True)
            validation.showErrorMessage = True
            validation.errorStyle = "stop"
            ws.add_data_validation(validation)
            validation.add(f"{col}2:{col}{ws.max_row}")
    for i in range(2, ws.max_row + 1):
        ws.row_dimensions[i].height = 180


def build(source: Path, reference: Path, target: Path):
    if target.exists():
        raise FileExistsError(f"Refusing to overwrite: {target}")
    before = {str(p): digest(p) for p in (source, reference)}
    original = load_workbook(source)
    old = original["Review"]
    headers = [c.value for c in old[1]]
    rows = [dict(zip(headers, row)) for row in old.iter_rows(min_row=2, values_only=True)]
    ref = json.loads(reference.read_text(encoding="utf-8"))
    reference_rows = {r["seed_id"]: r for r in ref["rows"]}
    if len(rows) != 20 or len({r["seed_id"] for r in rows}) != 20:
        raise ValueError("Expected 20 unique review rows")
    if set(reference_rows) != {r["seed_id"] for r in rows}:
        raise ValueError("Reference ID mismatch")
    evidence_columns = headers[:headers.index("factual_fidelity")]
    for row in rows:
        for key in evidence_columns:
            if (row[key] or "") != (reference_rows[row["seed_id"]].get(key) or ""):
                raise ValueError(f"Source evidence changed: {row['seed_id']} {key}")
    tail = headers[headers.index("review_decision"):]
    columns = evidence_columns + DIMENSIONS + tail
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.append(columns)
    for row in rows:
        # Preserve earlier answers, but changed rubric requires fresh completion.
        ws.append(["pending" if k == "review_status" else row.get(k) for k in columns])
    style_table(ws, CHOICES, DIMENSIONS + tail)
    ws.freeze_panes = "F2"

    guide = wb.create_sheet("Guide")
    guide.append(["欄位／步驟", "說明"])
    guide.append(["用途", "探索性 smoke review，不是 human gold、privacy clearance 或正式 benchmark 放行。沒有總分，不加總七個面向。"])
    guide.append(["如何閱讀", "先看 grounding，再看 generated dialogue；兩欄分開。長內容超過列高時可展開公式列，或用同目錄 review_readable.html 看完整對話。灰欄不改，黃欄填答。"])
    for key, description in DESCRIPTIONS.items():
        guide.append([key, description])
    guide.append(["選項", "pass＝未發現該面向的實質問題；fail＝有具體反例；uncertain＝資訊不足或界線需討論。不確定不能當 pass。"])
    guide.append(["問題記錄", "problem_evidence 用『面向：原文片段／理由』逐項記錄；label_consistency 加註 source_label_conflict / generated_label_drift / insufficient_evidence。多個面向可同時 fail。"])
    guide.append(["review_decision", "usable 僅限七項全 pass；needs_revision＝可修；unusable＝核心不可用；uncertain＝待釐清。非 usable 必須填 problem_evidence；needs_revision 再填 revision_suggestion。"])
    guide.append(["完成方式", "填 reviewer_kind、reviewer_id、UTC ISO 時間（如 2026-09-21T10:00:00Z），七項與處置均完成後才設 reviewed。AI 評審標 llm，不可冒充 human。"])
    guide.append(["原有評分", "若來源已有評分則保留，但全部重設 pending：新增面向與新版定義需再次確認。本次不自動評分。"])
    guide.append(["Batch_review", "看完 20 筆後填一次。重複用語只是線索，要引用 seed_id 與片段；不要求把 20 筆記在腦中。未取得 conditioning 設定時應標 uncertain，不凭對話反推設定。"])
    guide.append(["自動檢查", "原有 deterministic_flags 只是查核提示，不是判決；輪數/schema/regex 另看 technical_check_report.json，不能代替七個面向。"])
    guide.append(["統計相容性", "舊 notebook summarize_review 只懂四面向，不可拿來彙總本新版。填完後需以七面向驗證完成性、逐面向計數並區分 human/llm；pending 不算通過。"])
    style_table(guide)
    guide.column_dimensions["B"].width = 110
    for i in range(2, guide.max_row + 1):
        guide.row_dimensions[i].height = 65

    batch = wb.create_sheet("Batch_review")
    bh = ["aspect", "question", "assessment", "evidence_seed_ids_and_excerpts",
          "revision_suggestion", "reviewer_kind", "reviewer_id", "reviewed_utc", "review_status"]
    batch.append(bh)
    questions = {
        "opening_and_phrase_diversity": "開場、安慰語與句型是否過度重複？引用不同案例，不以單一詞頻定罪。",
        "interaction_strategy_diversity": "不同爭議是否仍套同一套追問與處理流程？追問是否配合個案？",
        "mood_conditioning": "先逐筆核對實際 customer_mood 設定，再比较不同 mood 的表現；設定不可得標 uncertain。",
        "length_and_information_gain": "對照實際 conversation_length 設定與訊息數；較長對話是否增加有用資訊，而非填充重複內容？",
        "recurring_failure_modes": "綜合七面向：哪些問題反覆出現？應修改 seed、prompt 或格式檢查哪個部分？",
    }
    for key, question in questions.items():
        batch.append([key, question, None, None, None, None, None, None, "pending"])
    style_table(batch, {"assessment": "no_material_issue,issue_found,uncertain",
                       "reviewer_kind": CHOICES["reviewer_kind"], "review_status": CHOICES["review_status"]}, bh[2:])
    meta = wb.create_sheet("_lineage")
    meta.append(["schema", "expanded_generation_review_v02"])
    meta.append(["source_hashes", json.dumps(before, ensure_ascii=False)])
    meta.append(["source_lineage", json.dumps(ref["lineage"], ensure_ascii=False)])
    meta.append(["policy", "smoke_only; privacy_unverified; no_auto_acceptance; old_four_dimension_summary_incompatible"])
    meta.sheet_state = "hidden"
    original.close()
    # Exclusive file creation preserves even a concurrently created target.
    with target.open("xb") as handle:
        wb.save(handle)
    if before != {str(p): digest(p) for p in (source, reference)}:
        raise RuntimeError("Source changed during export")
    return target


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("review_dir", type=Path)
    parser.add_argument("--output-name", default="generation_review_20_v03.xlsx")
    args = parser.parse_args()
    print(build(args.review_dir / "generation_review_20.xlsx",
                args.review_dir / "review_reference.json", args.review_dir / args.output_name))
